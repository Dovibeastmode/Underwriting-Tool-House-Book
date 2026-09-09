import openpyxl, os, subprocess
from openpyxl.styles import PatternFill, Font

SRC9="/root/.claude/uploads/15c93af4-845a-5d84-bd2c-2f824a3823d2/c2b9a3ea-Underwriting_Tool_Combined_with_House_book_analytics.xlsx"
LAST=1068
HD="'Historical Data'!"
BLUE=PatternFill("solid", fgColor="FFCCE5FF")

DIMS=[("K","AR"),("L","AT"),("M","AS"),("N","AQ")]   # flag col -> match col
TIER_FLAGS={14:(1,1,1,0),15:(1,0,1,1),16:(1,1,0,1),17:(1,0,1,0),
            18:(1,1,0,0),19:(1,0,0,1),20:(1,0,0,0),21:(0,0,0,0)}

def mask(r):
    return "*".join(f"({HD}${m}$2:${m}${LAST}*${f}{r}+(1-${f}{r}))" for f,m in DIMS)

def unmerge(ws,coord):
    for rng in list(ws.merged_cells.ranges):
        if coord in rng: ws.unmerge_cells(str(rng))

def build(highlight=True):
    wb=openpyxl.load_workbook(SRC9)
    hs=wb["Historical Score"]; hd=wb["Historical Data"]; dl=wb["Deal"]
    touched=[]
    def put(ws,coord,val,note=None):
        unmerge(ws,coord)
        ws[coord]=val
        if highlight: ws[coord].fill=BLUE
        touched.append((ws.title,coord,note))

    # 1 - flag table
    for col,h in zip("KLMN",["req NR","req ISO","req Ind","req Pos"]):
        put(hs,f"{col}13",h,"header")
    put(hs,"O13","<- tier definition: 1 = must match, 0 = ignore","note")
    for r,fl in TIER_FLAGS.items():
        for col,v in zip("KLMN",fl): put(hs,f"{col}{r}",v,"flag")

    # 2 - generic mask, identical in all 8 tier rows
    for r in TIER_FLAGS:
        M=mask(r)
        put(hs,f"C{r}",f"=SUMPRODUCT({HD}$AP$2:$AP${LAST}*{M})","seasoned deals")
        put(hs,f"D{r}",f"=SUMPRODUCT({HD}$AP$2:$AP${LAST}*{M}*{HD}$N$2:$N${LAST})","funded $")
        put(hs,f"F{r}",f'=IF(D{r}=0,"",SUMPRODUCT({HD}$AP$2:$AP${LAST}*{M}*{HD}$AK$2:$AK${LAST})/D{r})',"PLR")
        put(hs,f"G{r}",f'=IF(D{r}=0,"",SUMPRODUCT({HD}$AP$2:$AP${LAST}*{M}*{HD}$AJ$2:$AJ${LAST})/D{r})',"ROF")

    # 3 - selection
    put(hs,"D29",'=IF($C$5="","",MATCH("Yes",$E$14:$E$21,0))',"IFERROR gone; blank-deal guard added")
    put(hs,"E29",'row # (read by Historical Data)',"label")
    for col in "KLMN":
        put(hs,f"{col}29",f'=IF($D$29="","",INDEX({col}$14:{col}$21,$D$29))',"selected tier flag")
    put(hs,"O29","<- flags of the selected tier (read by Historical Data BA)","note")

    # 4 - one cell owns the no-evidence question
    put(hs,"D30",'=COUNTA($B$14:$B$21)',"fallback row #")
    put(hs,"E30",'fallback row # (whole seasoned book)',"label")
    put(hs,"B37","HAS EVIDENCE?","label")
    put(hs,"C37",'=IF(OR($D$29="",$D$29=$D$30),"NO","YES")',"THE guard cell")
    put(hs,"D37","<- every guard downstream reads this, not a hard-coded row number","note")

    # 5 - result block reads the guard
    put(hs,"C29",'=IF($D$29="","- no deal entered -",INDEX($B$14:$B$21,$D$29))',"blank-deal text")
    for c,src,nm in [("C30","$C$14:$C$21","seasoned deals"),("C31","$F$14:$F$21","PLR"),
                     ("C32","$G$14:$G$21","ROF"),("C33","$H$14:$H$21","index")]:
        put(hs,c,f'=IF($C$37="NO","",INDEX({src},$D$29))',nm)
    put(hs,"C34",'=IF(OR($C$37="NO",$C$33=""),"Neutral",IF($C$33<=Controls!$B$29,"Supportive",'
                 'IF($C$33>Controls!$B$30,"Negative","Neutral")))',"outcome")

    # 6 - retire the hardcoded combo columns
    for col in ["AU","AV","AW","AX","AY","AZ"]:
        for r in range(1,LAST+1):
            unmerge(hd,f"{col}{r}"); hd[f"{col}{r}"]=None
    put(hd,"AU1","(retired - combos now defined by the flag table on Historical Score)","6 columns cleared")

    # 7 - BA: mask against the selected tier, no CHOOSE
    for r in range(2,LAST+1):
        prod="*".join(f"({m}{r}*'Historical Score'!${f}$29+1-'Historical Score'!${f}$29)" for f,m in DIMS)
        c=f"BA{r}"; unmerge(hd,c); hd[c]=f"=IF(AD{r}=1,{prod},0)"
        if highlight and r<=LAST: hd[c].fill=BLUE
    touched.append(("Historical Data",f"BA2:BA{LAST}","CHOOSE replaced by the mask"))

    # 8 - Deal guards (outputs preserved: 0 in C56, text in F51)
    put(dl,"F51",'=IF(\'Historical Score\'!$C$37="NO","NO DEALS",'
                 'IF(\'Historical Score\'!$C$33<=Controls!$B$29,"GREEN",'
                 'IF(\'Historical Score\'!$C$33>Controls!$B$30,"RED","YELLOW")))',"guard -> C37")
    put(dl,"C56",'=IF(\'Historical Score\'!$C$37="NO",0,\'Historical Score\'!$C$30)',"guard -> C37")
    return wb,touched
