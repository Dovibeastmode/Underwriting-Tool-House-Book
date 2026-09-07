"""v4: no Decision Summary / Offer Engine / Checks; three-outcome house-book modifier; Section F outline; colour-blind Controls.
Usage: build_model_v4.py OUT.xlsx [demo]"""
import openpyxl, datetime as dt, copy, sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter as L
from openpyxl.cell.cell import MergedCell
from openpyxl.workbook.properties import CalcProperties
ROOT='/home/user/Underwriting-Tool-House-Book'
UW=f'{ROOT}/originals/Updated_Underwriting_Tool (v2 upload).xlsx'
FULL=f'{ROOT}/originals/House Book Analytics (2).xlsx'
SAMPLE=f'{ROOT}/originals/Copy_of_House_Book_Analytics (sample).xlsx'
OUT=sys.argv[1]; DEMO=len(sys.argv)>2 and sys.argv[2]=='demo'

DARK=PatternFill('solid',fgColor='1F2937'); MID=PatternFill('solid',fgColor='4B5563'); GREY=PatternFill('solid',fgColor='E5E7EB'); INPUT=PatternFill('solid',fgColor='E5E7EB')
WHITE=Font(name='Calibri',size=11,bold=True,color='FFFFFF'); SUB=Font(name='Calibri',size=9,bold=True,color='FFFFFF')
B=Font(name='Calibri',size=11,bold=True); N=Font(name='Calibri',size=11); NOTE=Font(name='Calibri',size=9,color='374151'); LINK=Font(name='Calibri',size=11,italic=True,color='374151'); INP=Font(name='Calibri',size=11,bold=True,color='000000')
thin=Side(style='thin',color='9CA3AF'); med=Side(style='medium',color='374151')
BOX=Border(left=thin,right=thin,top=thin,bottom=thin); IBOX=Border(left=med,right=med,top=med,bottom=med)
CEN=Alignment(horizontal='center',vertical='center'); LEFT=Alignment(horizontal='left',vertical='center'); WRAP=Alignment(horizontal='left',vertical='top',wrap_text=True)
G_FILL=PatternFill('solid',fgColor='00B050'); Y_FILL=PatternFill('solid',fgColor='FFFF00'); R_FILL=PatternFill('solid',fgColor='FF0000')
def rag_rules(ws,rng):
    for v,f,fc in (('GREEN',G_FILL,'FFFFFF'),('YELLOW',Y_FILL,'000000'),('RED',R_FILL,'FFFFFF'),('KNOCKOUT',R_FILL,'FFFFFF')):
        ws.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=[f'"{v}"'],fill=f,font=Font(color=fc,bold=True)))
def dec_rules(ws,rng):
    for v,f,fc in (('APPROVE',G_FILL,'FFFFFF'),('CONDITIONAL',Y_FILL,'000000'),('DECLINE',R_FILL,'FFFFFF'),('MANUAL REVIEW',PatternFill('solid',fgColor='F97316'),'FFFFFF')):
        ws.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=[f'"{v}"'],fill=f,font=Font(color=fc,bold=True)))
def chk_rules(ws,rng):
    for v,f,fc in (('OK',G_FILL,'FFFFFF'),('WARN',Y_FILL,'000000'),('FAIL',R_FILL,'FFFFFF')):
        ws.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=[f'"{v}"'],fill=f,font=Font(color=fc,bold=True)))
def hdr(ws,cell,text,span=None):
    ws[cell]=text; ws[cell].font=WHITE; ws[cell].fill=DARK; ws[cell].alignment=LEFT
    if span: ws.merge_cells(span)
def sub(ws,row,labels,start=1):
    for i,t in enumerate(labels):
        c=ws.cell(row,start+i,t); c.font=SUB; c.fill=MID; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=BOX
def put(ws,addr,val,font=N,fmt=None,align=CEN,border=BOX,fill=None):
    c=ws[addr]; c.value=val; c.font=font; c.alignment=align; c.border=border
    if font is INP and fill is None: fill=INPUT
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    return c
def label(ws,addr,text,font=N):
    c=ws[addr]; c.value=text; c.font=font; c.alignment=LEFT; c.border=BOX; return c
def note(ws,addr,text):
    c=ws[addr]; c.value=text; c.font=NOTE; c.alignment=LEFT; return c
def widths(ws,d):
    for k,v in d.items(): ws.column_dimensions[k].width=v

# ---------- house-book data ----------
hv=openpyxl.load_workbook(FULL,data_only=True); ref=hv['Ref']
sic=[(ref.cell(r,1).value,ref.cell(r,2).value,ref.cell(r,3).value) for r in range(4,68)]
hols=[ref.cell(r,6).value for r in range(4,45) if ref.cell(r,6).value]
fullraw=[list(r) for r in hv['Data'].iter_rows(min_row=2,max_row=1068,max_col=27,values_only=True) if r[0] is not None]
isos=sorted({r[12] for r in fullraw if r[12]},key=str.lower); states=sorted({r[11] for r in fullraw if r[11]})
statuses=[('Closed','No','Yes','Yes','No'),('Open','No','No','No','Yes'),('Legal','Yes','No','No','No'),('In House Collection','Yes','No','No','No'),('Lowered Payments','No','No','No','No'),('Written-Off','Yes','Yes','No','No'),('Bankruptcy','Yes','Yes','No','No'),('Settlement Agreement','Yes','Yes','No','No'),('Closed/NSF','No','Yes','Yes','No'),('Closed/NSF Unpaid','No','Yes','Yes','No'),('At-Risk','Yes','No','No','No')]
industries=sorted({s[1] for s in sic})+['No SIC','Unclassified']
raw=[list(r) for r in openpyxl.load_workbook(SAMPLE,data_only=True)['Data'].iter_rows(min_row=2,max_row=9,max_col=27,values_only=True)] if DEMO else fullraw
NR=1+len(raw)  # last data row
HDR="'Historical Data'!"; HS="'Historical Score'!"; OE="'Offer Engine'!"; DS="'Decision Summary'!"; C='Controls!'
def hr(col): return f"{HDR}${col}$2:${col}${NR}"

# ---------- Deal sheet from the template ----------
wb=openpyxl.load_workbook(UW)
for n in wb.sheetnames:
    if n!='New Template': del wb[n]
ws=wb['New Template']; ws.title='Deal'
def shift(ws,first,last,n):
    merges=[m for m in list(ws.merged_cells.ranges) if first<=m.min_row<=last]
    for m in merges: ws.unmerge_cells(str(m))
    for r in range(last,first-1,-1):
        for c in range(2,14):
            ws.cell(r+n,c)._style=copy.copy(ws.cell(r,c)._style); ws.cell(r+n,c).value=None
        if ws.row_dimensions[r].height: ws.row_dimensions[r+n].height=ws.row_dimensions[r].height
    for m in merges: ws.merge_cells(start_row=m.min_row+n,start_column=m.min_col,end_row=m.max_row+n,end_column=m.max_col)
    for r in range(first,first+n):
        for c in range(2,14):
            if not isinstance(ws.cell(r,c),MergedCell): ws.cell(r,c).value=None
shift(ws,45,63,1)          # credit-score row
for c in range(2,14): ws.cell(45,c)._style=copy.copy(ws.cell(44,c)._style)
shift(ws,49,64,10)         # C2 house-book indicator block
for r in range(49,59):     # style C2 rows like the metric rows
    for c in range(2,14): ws.cell(r,c)._style=copy.copy(ws.cell(38 if r>=51 else (36 if r==49 else 37),c)._style)
for m in list(ws.merged_cells.ranges):
    if m.min_row in (11,18) and m.min_col>=10: ws.unmerge_cells(str(m))
for r in range(11,19):
    for c in range(10,14):
        if not isinstance(ws.cell(r,c),MergedCell): ws.cell(r,c).value=None
for r in range(59,75):
    for c in range(2,14):
        if not isinstance(ws.cell(r,c),MergedCell): ws.cell(r,c).value=None
ws['E35']=None
ws.conditional_formatting=ConditionalFormattingList()
ws.data_validations.dataValidation=[dv for dv in ws.data_validations.dataValidation if str(dv.sqref) in ('E28:E33','B12:B23','E6')]

# rows (after shifts): metrics 38-45, info 46-47, C2 49-57, D 59-62, E 64-73, F 75+
ws['C2']='=C62'
ws['D2']='=IF(C62="","",C60&" red flag(s) vs limit "&C61&"   |   "&COUNTIF(F38:F45,"GREEN")&"G / "&COUNTIF(F38:F45,"YELLOW")&"Y / "&COUNTIF(F38:F45,"RED")&"R   |   score "&TEXT(C65,"0.0")&"/"&D65&" = "&C66'
ws['F6']=f'={C}$B$15'; ws['F6'].font=Font(name='Calibri',size=12,bold=True,color='008000')
ws['J6']='=IF($E$6="Weekly","Weekly remittance","Daily remittance")'
ws['K6']='=IF(OR($E$6="",C8="",C24=0),"",IF($E$6="Weekly",D8,C8))'
ws['L6']='=IF(OR(L7="",C24=0),"",L7*C24/IF($E$6="Weekly",$L$8,$L$9))'
ws['L7']='=IF(OR($E$6="",C8="",C24=0),"",ROUND(IF($E$6="Weekly",D8*$L$8,C8*$L$9)/C24*100,2)/100)'
ws['K8']=f'={C}$B$28'; ws['K9']=f'={C}$B$29'; ws['L8']=f'={C}$B$30'; ws['L9']=f'={C}$B$31'
for a in ['K8','K9','L8','L9']: ws[a].font=LINK
ws['E8']='=IF(C8="","",C8*$K$9)'; ws['F8']='=IF(OR(D6="",D6=0),"",IF(E6="Weekly",D6,D6/5)/($K$9/5))'
# Deal profile J11:M16
hdr(ws,'J11','DEAL PROFILE  (house-book match inputs)','J11:M11')
prof=[('ISO','K12','dropdown'),('Industry (SIC major group)','K13','dropdown'),('New deal or renewal','K14','REQUIRED for a cohort'),('Credit score (FICO)','K15','number')]
for i,(t,a,n_) in enumerate(prof):
    r=12+i; c=ws.cell(r,10,t); c.font=Font(name='Calibri',size=9,bold=True,color='FFFFFF'); c.fill=MID; c.alignment=LEFT; c.border=BOX
    cell=ws[a]; cell.border=IBOX; cell.alignment=CEN; cell.font=Font(name='Calibri',size=11,bold=True); cell.fill=INPUT
ws['L12']=f'=IF(K12="","",COUNTIF({hr("K")},K12)&" deals in book")'
ws['L13']=f'=IF(K13="","",COUNTIF({hr("Y")},K13)&" deals in book")'
ws['L14']='=IF(K14="","required for history match","")'
ws['L15']=f'=IF(K15="","not scored",IF(K15>={C}$B$12,"GREEN",IF(K15>={C}$C$12,"YELLOW","RED")))'
for r in range(12,16): ws.merge_cells(f'L{r}:M{r}'); ws[f'L{r}'].font=NOTE
note(ws,'J16','Position for matching = C42 (existing positions + 1). Bands and thresholds: Controls.')
for a,src in [('K12','=Ref!$E$2:$E$'+str(1+len(isos))),('K13','=Ref!$G$2:$G$'+str(1+len(industries))),('K14','"New Deal,Renewal"')]:
    dv=DataValidation(type='list',formula1=src,allow_blank=True); dv.add(a); ws.add_data_validation(dv)
for r in range(28,34):
    ws[f'F{r}']=f'=IF(D{r}<>"",IF(E{r}="Daily",D{r},IF(E{r}="Weekly",D{r}/5,D{r}/$K$9)),IF(C{r}="","",C{r}*{C}$B$33/({C}$B$23*$K$9)))'
    ws[f'H{r}']=f'=IF(OR(C{r}="",D{r}=""),"",ROUND(C{r}*{C}$B$33/D{r},0))'
    ws[f'I{r}']=f'=IF(H{r}="","",H{r}/IF(E{r}="Daily",$K$9,IF(E{r}="Weekly",$K$9/5,1)))'
    ws[f'J{r}']=f'=IF(OR(I{r}="",G{r}=""),"",MAX(0,I{r}-(TODAY()-G{r})/{C}$B$32))'
ws['C34']='=SUMPRODUCT(((C28:C33<>"")+(D28:D33<>"")>0)*(E28:E33<>"Monthly"))'
ws['F34']='=E34*$K$9'
# C. metrics
ws['H37']='Sizing line'; ws['H37']._style=copy.copy(ws['G37']._style)
metrics=[(38,'Total LVG (all debt / rev)',5,'=IF(C24=0,"",(E34+IF(C8="",0,C8))*$K$9/C24)','0.00%'),(39,'New holdback SP%',6,'=IF(OR(C24=0,E8=""),"",E8/C24)','0.00%'),(40,'Daily debt / balance',7,'=IF(D24=0,"",(E34+IF(C8="",0,C8))/D24)','0.00%'),(41,'Fund / revenue',8,'=IF(OR(C24=0,B6=""),"",B6/C24)','0.00'),(42,'# positions (incl. new)',9,'=IF(C24=0,"",C34+1)','0'),(43,'Revenue trend (avg MoM)',10,'=G24','0.0%'),(44,'Neg days + NSFs (total)',11,'=E24+F24','0'),(45,'Credit score (FICO)',12,'=IF(K15="","",K15)','0')]
for r,t,cr,f,fmt in metrics:
    ws[f'B{r}']=t; ws[f'C{r}']=f; ws[f'C{r}'].number_format=fmt
    for col,cc in (('D','B'),('E','C'),('G','D')):
        ws[f'{col}{r}']=f'={C}${cc}${cr}'; ws[f'{col}{r}'].font=LINK; ws[f'{col}{r}'].number_format='0.0%' if r in (38,39,40,43) else ('0.00' if r==41 else '0')
    if r==45: ws[f'F{r}']=f'=IF(C{r}="","",IF(AND(G{r}<>0,C{r}<=G{r}),"KNOCKOUT",IF(C{r}>=D{r},"GREEN",IF(C{r}<E{r},"RED","YELLOW"))))'
    elif r==43: ws['F43']='=IF(C24=0,"",IF(NOT(ISNUMBER(C43)),"YELLOW",IF(C43<=G43,"KNOCKOUT",IF(C43>=D43,"GREEN",IF(C43<=E43,"RED","YELLOW")))))'
    else: ws[f'F{r}']=f'=IF(C{r}="","",IF(C{r}>=G{r},"KNOCKOUT",IF(C{r}<=D{r},"GREEN",IF(C{r}>=E{r},"RED","YELLOW"))))'
    if r<=41:  # sizing line = green + flex x (red - green)
        ws[f'H{r}']=f'=D{r}+{HS}$C$36*(E{r}-D{r})'; ws[f'H{r}']._style=copy.copy(ws[f'C{r}']._style); ws[f'H{r}'].number_format=ws[f'D{r}'].number_format; ws[f'H{r}'].font=Font(name='Calibri',size=11,bold=True,color='000000')
ws['B46']='Existing leverage (before new)'; ws['C46']='=IF(C24=0,"",F34/C24)'; ws['D46']='-'; ws['E46']='-'; ws['F46']='INFO'
ws['B47']='Balance / revenue'; ws['C47']='=IF(C24=0,"",D24/C24)'; ws['D47']='-'; ws['E47']='-'; ws['F47']='INFO'
ws['C46'].number_format='0.0%'; ws['C47'].number_format='0.0%'
# C2. house-book indicators
ws['B49']='C2.  HOUSE-BOOK INDICATORS   (set the sizing line + modifier; NOT counted in the cash-flow score)'
for h,t in zip('BCDEFGH',['Cohort','Principal lost rate','Green <=','Red >','Status','Seasoned n','Return on funded']): ws[f'{h}50']=t
ws['H50']._style=copy.copy(ws['G50']._style)
book=f'{HS}$C$10'
def ind(r,lbl,hsrow):
    ws[f'B{r}']=lbl; ws[f'C{r}']=f'=IF({HS}$F${hsrow}="","",{HS}$F${hsrow})'; ws[f'C{r}'].number_format='0.0%'
    ws[f'H{r}']=f'=IF({HS}$G${hsrow}="","",{HS}$G${hsrow})'; ws[f'H{r}']._style=copy.copy(ws[f'C{r}']._style); ws[f'H{r}'].number_format='0.0%'
    ws[f'D{r}']=f'={book}*{C}$B$49'; ws[f'E{r}']=f'={book}*{C}$B$50'; ws[f'D{r}'].number_format=ws[f'E{r}'].number_format='0.0%'; ws[f'D{r}'].font=ws[f'E{r}'].font=LINK
    ws[f'F{r}']=f'=IF({HS}$C${hsrow}=0,"NO DEALS",IF({HS}$F${hsrow}="","TOO FEW",IF({HS}$H${hsrow}<={C}$B$49,"GREEN",IF({HS}$H${hsrow}<={C}$B$50,"YELLOW","RED"))))'
    ws[f'G{r}']=f'={HS}$C${hsrow}'; ws[f'G{r}'].number_format='0'
ws['B51']=f'="Primary cohort: "&{HS}$C$30'; ws['C51']=f'=IF({HS}$C$29=7,"",{HS}$C$32)'; ws['C51'].number_format='0.0%'
ws['H51']=f'=IF({HS}$C$29=7,"",{HS}$C$33)'; ws['H51']._style=copy.copy(ws['C51']._style); ws['H51'].number_format='0.0%'
ws['D51']=f'={book}*{C}$B$49'; ws['E51']=f'={book}*{C}$B$50'; ws['D51'].number_format=ws['E51'].number_format='0.0%'; ws['D51'].font=ws['E51'].font=LINK
ws['F51']=f'=IF({HS}$C$29=7,"NO EVIDENCE",IF({HS}$C$34<={C}$B$49,"GREEN",IF({HS}$C$34<={C}$B$50,"YELLOW","RED")))'; ws['G51']=f'={HS}$C$31'
ind(52,'ISO',25); ind(53,'Industry',24); ind(54,'New deal / renewal',23); ind(55,'Position',22)
ws['B56']='Seasoned deals behind the primary cohort'; ws['C56']=f'=IF({HS}$C$29=7,0,INDEX({HS}$C$14:$C$20,{HS}$C$29))'; ws['C56'].number_format='0'
ws['D56']=f'={C}$B$44'; ws['E56']=f'={C}$B$40'; ws['D56'].font=ws['E56'].font=LINK; ws['D56'].number_format=ws['E56'].number_format='0'
ws['F56']='=IF(C56>=D56,"GREEN",IF(C56>=E56,"YELLOW","RED"))'
ws['G56']=f'="house book: "&{HS}$C$35&"  |  sizing flex "&TEXT({HS}$C$36,"0%")'; ws['G56'].alignment=LEFT
ws.merge_cells('G56:H56')
for c in range(2,9): ws.cell(57,c).value=None
rag_rules(ws,'F38:F45'); rag_rules(ws,'F51:F56'); dec_rules(ws,'C2'); dec_rules(ws,'C62')
# D. decision
ws['B59']='D.  DECISION  (cash-flow layer, on the requested offer)'
ws['B60']='Red flags'; ws['C60']='=COUNTIF(F38:F45,"RED")+COUNTIF(F38:F45,"KNOCKOUT")'; ws['D60']='Knockouts'; ws['E60']='=COUNTIF(F38:F45,"KNOCKOUT")'
ws['B61']='Decline if red flags >=  (Controls)'; ws['C61']=f'={C}$B$13'; ws['D61']='Minimum fund $'; ws['E61']=f'={C}$B$14'
ws['C61'].font=Font(name='Calibri',size=12,bold=True,color='008000'); ws['E61'].font=Font(name='Calibri',size=12,bold=True,color='008000')
ws['B62']='DECISION'
ws['C62']='=IF(C24=0,"",IF(OR(E60>0,IF(ISNUMBER(C73),C73,0)<E61),"DECLINE",IF(OR(B6="",B6=0),"",IF(C60>=C61,"DECLINE",IF(COUNTIF(F38:F45,"GREEN")=D65,"APPROVE","CONDITIONAL")))))'
ws['D62']='=IF(C62="DECLINE",IF(E60>0,"KNOCKOUT: "&INDEX(B38:B45,MATCH("KNOCKOUT",F38:F45,0))&" is past its knockout line.",IF(C60>=C61,"Red flags reached the threshold.","Max fund "&IF(ISNUMBER(C73),TEXT(C73,"$#,##0"),"$0")&" is below your minimum "&TEXT(E61,"$#,##0")&".")),IF(C62="APPROVE","All "&D65&" decision metrics green.",IF(C62="","Enter offer + bank data.","Fundable but not clean - review yellows/reds above.")))'
# E. sizing at the SIZING LINE (= green line when the modifier is neutral)
ws['B64']='E.  DEAL QUALITY  &  MAXIMUM OFFER at the sizing line  (green line unless the house book is supportive)'
ws['B65']='Score (G=1, Y=0.5, R=0)'; ws['C65']='=COUNTIF(F38:F45,"GREEN")+0.5*COUNTIF(F38:F45,"YELLOW")'; ws['D65']='=SUMPRODUCT(--(F38:F45<>""))'; ws['E65']='scored metrics'
ws['B66']='Tier  (score scaled to 7 for the bands)'; ws['C66']=f'=IF(C24=0,"",IF(D67>={C}$B$25,"Strong",IF(D67>={C}$C$25,"Solid",IF(D67>={C}$D$25,"Moderate","Weak"))))'
ws['B67']='Max daily by SP% sizing line'; ws['C67']='=IF(C24=0,"",H39*C24/$K$9)'; ws['D67']='=IF(D65=0,0,C65*7/D65)'; ws['E67']='scaled score /7'
ws['B68']='Max daily by Total LVG sizing line'; ws['C68']='=IF(C24=0,"",(H38*C24-F34)/$K$9)'
ws['B69']='Max daily by Daily/Balance sizing line'; ws['C69']='=IF(OR(C24=0,D24=0),"",H40*D24-E34)'
ws['B70']='MAX DAILY  (weakest link)'; ws['C70']='=IF(C24=0,"",MAX(0,MIN(C67:C69)))'
ws['D70']='=IF(C70="","",IF(C70=0,"no clean room - existing debt is over the sizing line",INDEX({"SP% holdback";"Total LVG";"Daily/Balance"},MATCH(C70,C67:C69,0))&" is the limit"))'
ws['B71']='Weekly / monthly equivalent'; ws['C71']='=IF(OR(C70="",C70=0),"",C70*5)'; ws['D71']='=IF(OR(C70="",C70=0),"",C70*$K$9)'
ws['B72']='Term (daily pmts) / factor from bands'
ws['C72']=f'=IF(OR(C70="",C70=0,C65=""),"",MAX(20,MIN(INDEX({C}$B$19:$B$22,MATCH(D67,{C}$A$19:$A$22,1)),FLOOR({C}$B$23*$K$9,5),FLOOR(H41*C24*D72/C70,5))))'
ws['D72']=f'=IF(OR(C65="",C24=0),"",INDEX({C}$C$19:$C$22,MATCH(D67,{C}$A$19:$A$22,1)))'
ws['B73']='MAX FUND at the sizing line'
ws['C73']='=IF(COUNTIF(F42:F45,"KNOCKOUT")>0,"-",IF(OR(C70="",C70=0,C72="",D72=""),"",FLOOR(C70*C72/D72,500)))'
ws['D73']='=IF(C73="-","KNOCKOUT on the merchant - do not fund.",IF(C73="","","payback "&TEXT(C73*D72,"$#,##0")&" over "&C72&" daily pmts ("&TEXT(C72/5,"0")&" wks)"))'
for a in ['C67','C68','C69','C70']: ws[a].number_format='"$"#,##0.00'
ws['D67'].number_format='0.0'; ws['C71'].number_format=ws['D71'].number_format='"$"#,##0'; ws['C72'].number_format='0'; ws['D72'].number_format='0.00'; ws['C73'].number_format='"$"#,##0'
# F. final offer - OUTLINE (planned, no formulas yet)
PLAN=PatternFill('solid',fgColor='F3F4F6'); LOGIC=Font(name='Calibri',size=10,italic=True,color='374151')
r0=75
hdr(ws,f'B{r0}','F.  FINAL OFFER  (PLANNED - outline only)   =   max at the sizing line (section E)  ->  house-book outcome applied',f'B{r0}:H{r0}')
sub(ws,r0+1,['Item','Value','How it will be computed'],start=2); ws.merge_cells(f'D{r0+1}:H{r0+1}')
plan=[('FINAL DECISION','APPROVE / CONDITIONAL / DECLINE','DECLINE if D says DECLINE on the merchant side (knockout, red flags, below minimum). CONDITIONAL if any metric is YELLOW at the final offer or the house-book outcome is Negative. Else APPROVE.'),
      ('Recommended fund','$','Section E max fund (C73). If the house-book outcome is Negative: C73 x (1 - fund cut % on Controls). Floor to $500. Never above C73.'),
      ('Factor','1.40 / 1.45 / 1.49','Band factor D72 (unchanged by history for now).'),
      ('Term','weeks or daily pmts','Band term C72; shown in the frequency of E6.'),
      ('Payment','per week / per day','fund x factor / term (daily); x5 for weekly.'),
      ('Holdback % at final offer','%','payment x 21.655 / avg revenue, with GREEN / YELLOW status against Controls. Yellow is allowed only when the outcome is Supportive.'),
      ('Total LVG / daily debt-balance / fund-rev at final offer','%','same re-check for the other three offer-side metrics; RED must be impossible because the sizing line stops short of the red line.'),
      ('House-book outcome','Supportive / Neutral / Negative','from Historical Score C35'),
      ('Primary cohort and its numbers','text','tier, seasoned n, PLR, ROF, index - from Historical Score')]
for i,(t,val,how) in enumerate(plan):
    r=r0+2+i; c=ws.cell(r,2,t); c.font=B; c.alignment=Alignment(horizontal='right',vertical='center'); c.border=BOX
    c=ws.cell(r,3,val); c.font=NOTE; c.alignment=CEN; c.border=BOX
    c=ws.cell(r,4,'-> '+how); c.font=LOGIC; c.fill=PLAN; c.alignment=WRAP; c.border=BOX; ws.merge_cells(f'D{r}:H{r}'); ws.row_dimensions[r].height=30
note(ws,f'B{r0+2+len(plan)}','Layer 1 stays the ceiling: the sizing line never reaches the red line, knockouts are never removed, the final fund never exceeds C73.')
ws.column_dimensions['J'].width=27; ws.column_dimensions['K'].width=26; ws.column_dimensions['L'].width=16; ws.column_dimensions['M'].width=12; ws.column_dimensions['H'].width=13
ws.sheet_view.showGridLines=False

# ======================= Controls =======================
cs=wb.create_sheet('Controls'); widths(cs,{'A':52,'B':15,'C':15,'D':15,'E':15,'F':13,'G':13,'H':13,'I':60})
cs['A1']='CONTROLS  -  every editable threshold, tier, historical weight and rule'; cs['A1'].font=Font(name='Calibri',size=14,bold=True)
note(cs,'A2','EDITABLE cells = bold black text on grey with a thick border. Do not insert or delete rows; formulas reference fixed cells.')
hdr(cs,'A3','1. CASH-FLOW THRESHOLDS','A3:I3'); sub(cs,4,['Metric','Green (at or better)','Red (at or worse)','Knockout','Direction','','','','Note'])
thr=[('Total LVG (all debt / rev)',0.30,0.45,0.55,'higher is worse','0.0%',''),('New holdback SP%',0.18,0.25,0.35,'higher is worse','0.0%',''),('Daily debt / balance',0.18,0.30,0.35,'higher is worse','0.0%',''),('Fund / revenue',0.60,0.90,1.25,'higher is worse','0.00',''),('# positions (incl. new)',2,5,6,'higher is worse','0',''),('Revenue trend (avg MoM)',-0.05,-0.15,-0.30,'lower is worse','0.0%',''),('Neg days + NSFs (total)',1,6,10,'higher is worse','0','total across all months entered'),('Credit score (FICO)',650,630,0,'lower is worse','0','650+ green, 630-649 yellow, below 630 red; knockout 0 = off')]
for i,(t,g,r_,k,d,fmt,n_) in enumerate(thr):
    r=5+i; label(cs,f'A{r}',t); put(cs,f'B{r}',g,INP,fmt,border=IBOX); put(cs,f'C{r}',r_,INP,fmt,border=IBOX); put(cs,f'D{r}',k,INP,fmt,border=IBOX); put(cs,f'E{r}',d,N,align=LEFT); note(cs,f'I{r}',n_)
label(cs,'A13','Decline if red flags (RED + KNOCKOUT) >='); put(cs,'B13',2,INP,'0',border=IBOX)
label(cs,'A14','Minimum fund $'); put(cs,'B14',10000,INP,'"$"#,##0',border=IBOX)
label(cs,'A15','Underwriting fee %'); put(cs,'B15',0.05,INP,'0%',border=IBOX)
hdr(cs,'A17','2. PRICING BANDS','A17:I17'); sub(cs,18,['Score >=','Term cap (daily pmts)','Factor','','','','','','Note'])
for i,(s,t,f) in enumerate([(0,100,1.49),(5,130,1.45),(6,150,1.45),(7,195,1.40)]):
    r=19+i; put(cs,f'A{r}',s,INP,'0.0',border=IBOX); put(cs,f'B{r}',t,INP,'0',border=IBOX); put(cs,f'C{r}',f,INP,'0.00',border=IBOX)
label(cs,'A23','Max term (months)'); put(cs,'B23',9,INP,'0',border=IBOX)
label(cs,'A24','Approved factor tiers (low -> high)'); put(cs,'B24',1.40,INP,'0.00',border=IBOX); put(cs,'C24',1.45,INP,'0.00',border=IBOX); put(cs,'D24',1.49,INP,'0.00',border=IBOX)
label(cs,'A25','Tier cut-offs on 7-point scale: Strong / Solid / Moderate'); put(cs,'B25',6.5,INP,'0.0',border=IBOX); put(cs,'C25',5,INP,'0.0',border=IBOX); put(cs,'D25',3.5,INP,'0.0',border=IBOX)
hdr(cs,'A27','3. CONVERSION CONSTANTS','A27:I27')
for r,(t,v,fmt) in enumerate([('Weeks per month (ours)',4.331,'0.000'),('Business days per month (ours)',21.655,'0.000'),('MoneyBadger weeks per month',4,'0'),('MoneyBadger business days per month',21,'0'),('Calendar days per month',30.4375,'0.0000'),('Assumed factor on existing positions',1.45,'0.00')],start=28):
    label(cs,f'A{r}',t); put(cs,f'B{r}',v,INP,fmt,border=IBOX)
hdr(cs,'A35','4. HISTORICAL EVIDENCE SETTINGS','A35:I35')
hist=[('As-of date for seasoning (date of the CRM export)',dt.datetime(2026,8,25),'mm/dd/yyyy',''),('Window: funded on or after',dt.datetime(2024,11,1),'mm/dd/yyyy',''),('Window: funded on or before',dt.datetime(2026,8,25),'mm/dd/yyyy','12/31/2025 = settled book'),
      ('Seasoning multiple (age >= multiple x contracted term)',1.25,'0.00',''),('Minimum seasoned deals for a cohort to be PRIMARY',3 if DEMO else 10,'0','RECOMMENDED for the full book: 10' if DEMO else ''),('Minimum seasoned funded $ for a cohort to be PRIMARY',30000 if DEMO else 150000,'"$"#,##0','RECOMMENDED for the full book: $150,000' if DEMO else ''),
      ('(reserved) Credibility constant K - NOT USED in v3',30,'0','SIDE NOTE: thinking of adding a blend rate and a credibility factor later (small cohorts pulled toward the book by Z = n/(n+K)). Nothing reads this cell today.'),('Minimum seasoned deals before a rate is shown',2 if DEMO else 5,'0','RECOMMENDED for the full book: 5' if DEMO else ''),('Seasoned deals for the sample-size indicator to be GREEN',5 if DEMO else 30,'0','RECOMMENDED for the full book: 30' if DEMO else '')]
for i,(t,v,fmt,n_) in enumerate(hist):
    r=36+i; label(cs,f'A{r}',t); put(cs,f'B{r}',v,INP,fmt,border=IBOX); note(cs,f'I{r}',n_)
hdr(cs,'A47','5. HOUSE-BOOK OUTCOME   index = primary-cohort principal lost rate / seasoned-book principal lost rate  (1.00 = book)','A47:I47')
sub(cs,48,['Setting','Value','','','','','','','Note'])
label(cs,'A49','SUPPORTIVE when index at or below'); put(cs,'B49',0.85,INP,'0.00'); note(cs,'I49','cohort loses at most 85% of what the book loses')
label(cs,'A50','NEGATIVE when index above'); put(cs,'B50',1.15,INP,'0.00'); note(cs,'I50','between the two = NEUTRAL = the cash-flow tool exactly as built')
label(cs,'A51','Supportive: sizing flex (0 = green line, 1 = red line)'); put(cs,'B51',0.50,INP,'0%'); note(cs,'I51','RECOMMENDED 50%: size half-way to the red line, so metrics may go YELLOW, never RED. Keep below 100%.')
label(cs,'A52','Negative: fund cut %'); put(cs,'B52',0.20,INP,'0%'); note(cs,'I52','RECOMMENDED 20%: green-line max x (1 - cut). Section F, planned.')
note(cs,'A54','The C2 block on the Deal sheet uses B49 / B50 as its GREEN / RED lines for every house-book indicator. Outcome = Neutral whenever no cohort meets the minimums (tier 7).')
hdr(cs,'A57','6. STATUS MAPPING','A57:I57'); sub(cs,58,['Status','Default?','Resolved?','Closed?','Open?','Deals in file','','',''])
for i,s in enumerate(statuses):
    r=59+i; label(cs,f'A{r}',s[0])
    for j in range(4): put(cs,f'{L(2+j)}{r}',s[j+1],INP,border=IBOX)
    put(cs,f'F{r}',f'=COUNTIF({hr("D")},A{r})',N,'0')
cs.sheet_view.showGridLines=False

# ======================= Ref =======================
rf=wb.create_sheet('Ref'); widths(rf,{'A':8,'B':40,'C':32,'E':44,'G':40,'I':10,'K':14,'M':22})
for col,t in zip('ABCEGKM',['SIC2','Industry (SIC major group)','Division','ISO (house-book spelling)','Industry list','Fed holiday','Status list']):
    c=rf[f'{col}1']; c.value=t; c.font=SUB; c.fill=MID; c.alignment=CEN; c.border=BOX
for i,(a,b,c_) in enumerate(sic): rf.cell(2+i,1,a); rf.cell(2+i,2,b); rf.cell(2+i,3,c_)
for i,v in enumerate(isos): rf.cell(2+i,5,v)
for i,v in enumerate(industries): rf.cell(2+i,7,v)
for i,v in enumerate(hols): rf.cell(2+i,11,v).number_format='mm/dd/yyyy'
for i,s in enumerate(statuses): rf.cell(2+i,13,s[0])
HOL=f'Ref!$K$2:$K${1+len(hols)}'

# ======================= Historical Data =======================
hs=wb.create_sheet('Historical Data')
cols=['Contract ID','DBA','Legal name','Status','Start date','Last active','Position','SIC','SIC2','State','ISO','New or renewal','Factor','Advance','Commission','Payback','Collected','Balance','% paid','Exp dur (raw)','NSF count','Avg rev/mo','Holdback %','Biz start',
      'Industry','Pay frequency','Term (bus. days)','Age at as-of (bus. days)','Seasoned','In window','Default','Resolved','Closed family','Open','Status valid','Net cash','Principal lost','Adv / revenue','Has revenue','Days on book (bus.)','Renewal',
      'Eligible (seasoned & in window)','m: position','m: new/renewal','m: industry','m: ISO','T1 NR+Ind+Pos','T2 NR+ISO+Pos','T3 NR+Ind','T4 NR+ISO','T5 NR+Pos','T6 NR','T7 Book','Best tier','In primary cohort']
for i,t in enumerate(cols):
    c=hs.cell(1,1+i,t); c.font=SUB; c.fill=MID if i<24 else DARK; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=BOX
hs.row_dimensions[1].height=42
keep=[0,1,2,3,4,5,6,7,9,11,12,13,14,15,16,17,18,19,20,21,23,24,25,26]
st="Controls!$A$59:$E$69"
for i,r in enumerate(raw):
    rr=2+i
    for j,k in enumerate(keep):
        v=r[k]; v=str(v) if (k in (7,9) and v is not None) else v
        hs.cell(rr,1+j,v)
    f={'Y':f'=IF(I{rr}="","No SIC",IFERROR(VLOOKUP(I{rr},Ref!$A$2:$B$65,2,FALSE),"Unclassified"))','Z':f'=IF(T{rr}="","",IF(T{rr}<12,"Daily","Weekly"))',
       'AA':f'=IF(T{rr}="",0,IF(T{rr}<12,T{rr}*Controls!$B$29,T{rr}*5))','AB':f'=IF(E{rr}="",0,NETWORKDAYS(E{rr},Controls!$B$36,{HOL}))','AC':f'=IF(AND(AA{rr}>0,AB{rr}>=Controls!$B$39*AA{rr}),1,0)',
       'AD':f'=IF(AND(E{rr}<>"",E{rr}>=Controls!$B$37,E{rr}<=Controls!$B$38),1,0)','AE':f'=IF(IFERROR(VLOOKUP(D{rr},{st},2,FALSE),"No")="Yes",1,0)','AF':f'=IF(IFERROR(VLOOKUP(D{rr},{st},3,FALSE),"No")="Yes",1,0)',
       'AG':f'=IF(IFERROR(VLOOKUP(D{rr},{st},4,FALSE),"No")="Yes",1,0)','AH':f'=IF(IFERROR(VLOOKUP(D{rr},{st},5,FALSE),"No")="Yes",1,0)','AI':f'=IF(ISNUMBER(MATCH(D{rr},Controls!$A$59:$A$69,0)),1,0)',
       'AJ':f'=Q{rr}-N{rr}','AK':f'=IF(AE{rr}=1,MAX(0,N{rr}-Q{rr}),0)','AL':f'=IF(OR(V{rr}="",V{rr}=0),0,N{rr}/V{rr})','AM':f'=IF(OR(V{rr}="",V{rr}=0),0,1)','AN':f'=IF(OR(E{rr}="",F{rr}=""),0,NETWORKDAYS(E{rr},F{rr},{HOL}))','AO':f'=IF(L{rr}="Renewal",1,0)','AP':f'=AC{rr}*AD{rr}',
       'AQ':f"=IF(AND(G{rr}<>\"\",{HS}$C$6<>\"\"),IF(G{rr}={HS}$C$6,1,0),0)",'AR':f"=IF(AND(L{rr}<>\"\",{HS}$C$5<>\"\"),IF(L{rr}={HS}$C$5,1,0),0)",'AS':f"=IF(AND(Y{rr}<>\"\",{HS}$C$7<>\"\"),IF(Y{rr}={HS}$C$7,1,0),0)",'AT':f"=IF(AND(K{rr}<>\"\",{HS}$C$8<>\"\"),IF(K{rr}={HS}$C$8,1,0),0)",
       'AU':f'=AR{rr}*AS{rr}*AQ{rr}','AV':f'=AR{rr}*AT{rr}*AQ{rr}','AW':f'=AR{rr}*AS{rr}','AX':f'=AR{rr}*AT{rr}','AY':f'=AR{rr}*AQ{rr}','AZ':f'=AR{rr}','BA':'=1',
       'BB':f'=IF(AU{rr}=1,1,IF(AV{rr}=1,2,IF(AW{rr}=1,3,IF(AX{rr}=1,4,IF(AY{rr}=1,5,IF(AZ{rr}=1,6,7))))))','BC':f"=IF(AD{rr}=1,CHOOSE({HS}$C$29,AU{rr},AV{rr},AW{rr},AX{rr},AY{rr},AZ{rr},BA{rr}),0)"}
    for col,fx in f.items(): hs[f'{col}{rr}']=fx
    for col in ('E','F','X'): hs[f'{col}{rr}'].number_format='mm/dd/yyyy'
    for col in ('N','O','P','Q','R'): hs[f'{col}{rr}'].number_format='"$"#,##0'
    hs[f'S{rr}'].number_format='0.0%'; hs[f'W{rr}'].number_format='0.0%'
hs.freeze_panes='E2'; widths(hs,{'A':12,'B':26,'C':26,'D':18,'K':24,'Y':26})

# ======================= Historical Score =======================
sc=wb.create_sheet('Historical Score'); widths(sc,{'A':7,'B':38,'C':16,'D':16,'E':12,'F':16,'G':16,'H':12,'I':60})
sc['A1']='HISTORICAL SCORE  -  which cohort, what it lost, what it returned'; sc['A1'].font=Font(name='Calibri',size=14,bold=True)
note(sc,'A2','Rates use SEASONED deals only (age >= multiple x contracted term). Index = cohort principal lost rate / seasoned-book principal lost rate (1.00 = book). SIDE NOTE: a blend rate and a credibility factor (small cohorts pulled toward the book) are under consideration; not built.')
hdr(sc,'A4','DEAL BEING MATCHED','A4:C4')
for r,(t,f) in enumerate([('New deal or renewal','=IF(Deal!$K$14="","",Deal!$K$14)'),('Position (Deal C42)','=IF(Deal!$C$42="","",Deal!$C$42)'),('Industry','=IF(Deal!$K$13="","",Deal!$K$13)'),('ISO','=IF(Deal!$K$12="","",Deal!$K$12)')],start=5):
    label(sc,f'B{r}',t); put(sc,f'C{r}',f,B)
label(sc,'B10','Seasoned-book principal lost rate (baseline)'); put(sc,'C10','=IF($F$20="",0,$F$20)',B,'0.0%')
label(sc,'B11','Seasoned-book return on funded'); put(sc,'C11','=IF($G$20="",0,$G$20)',B,'0.0%')
sub(sc,13,['Tier','Cohort','Seasoned deals','Funded $ (seasoned)','Meets min?','Principal lost rate','Return on funded','Index','How'])
sc.row_dimensions[13].height=30
def cohort_row(r,tier,name,flag):
    F=hr(flag); E_=hr('AP')
    put(sc,f'A{r}',tier,B); label(sc,f'B{r}',name)
    put(sc,f'C{r}',f'=SUMPRODUCT({F}*{E_})',N,'0')
    put(sc,f'D{r}',f'=SUMPRODUCT({F}*{E_}*{hr("N")})',N,'"$"#,##0')
    put(sc,f'E{r}',f'=IF(AND(C{r}>=Controls!$B$40,D{r}>=Controls!$B$41),"Yes","No")',B)
    put(sc,f'F{r}',f'=IF(OR(C{r}<Controls!$B$43,D{r}=0),"",SUMPRODUCT({F}*{E_}*{hr("AK")})/D{r})',N,'0.0%')
    put(sc,f'G{r}',f'=IF(OR(C{r}<Controls!$B$43,D{r}=0),"",SUMPRODUCT({F}*{E_}*{hr("AJ")})/D{r})',N,'0.0%')
    put(sc,f'H{r}',f'=IF(OR(F{r}="",$C$10=0),"",F{r}/$C$10)',N,'0.00')
tiers=[(1,'T1: new/renewal + industry + position','AU'),(2,'T2: new/renewal + ISO + position','AV'),(3,'T3: new/renewal + industry','AW'),(4,'T4: new/renewal + ISO','AX'),(5,'T5: new/renewal + position','AY'),(6,'T6: new/renewal only','AZ'),(7,'T7: seasoned book (baseline)','BA')]
for t,nme,flag in tiers: cohort_row(13+t,t,nme,flag)
note(sc,'I14','Seasoned deals = count of matches that are seasoned and in the window. Funded = their advances. Meets min = both Controls minimums. PLR = principal lost / funded. ROF = (collected - advance) / funded, closed deals included. Index = PLR / book PLR.')
sub(sc,21,['','SINGLE-DIMENSION EVIDENCE  (feeds the C2 block and the largest-impact test; does not pick the cohort)','','','','','','',''])
for i,(nme,flag) in enumerate([('Position only','AQ'),('New/renewal only','AR'),('Industry only','AS'),('ISO only','AT')]): cohort_row(22+i,'S'+str(i+1),nme,flag)
for r in range(22,26): put(sc,f'I{r}',f'=IF(F{r}="",0,ABS(F{r}-$C$10))',NOTE,'0.0%')
put(sc,'I21','|PLR - book PLR|',SUB,fill=MID)
sc.conditional_formatting.add('E14:E25',CellIsRule(operator='equal',formula=['"Yes"'],fill=G_FILL,font=Font(color='FFFFFF',bold=True)))
sc.conditional_formatting.add('E14:E25',CellIsRule(operator='equal',formula=['"No"'],fill=Y_FILL))
hdr(sc,'A28','RESULT','A28:C28')
res=[(29,'Primary tier = the LOWEST tier number that meets both minimums','=IFERROR(MATCH("Yes",$E$14:$E$20,0),7)','0'),
     (30,'Primary cohort','=INDEX($B$14:$B$20,$C$29)',None),
     (31,'Seasoned deals in it','=INDEX($C$14:$C$20,$C$29)','0'),
     (32,'Its principal lost rate','=INDEX($F$14:$F$20,$C$29)','0.0%'),
     (33,'Its return on funded','=INDEX($G$14:$G$20,$C$29)','0.0%'),
     (34,'Index (its PLR / book PLR)','=INDEX($H$14:$H$20,$C$29)','0.00'),
     (35,'HOUSE-BOOK OUTCOME','=IF(OR(C29=7,C34=""),"Neutral",IF(C34<=Controls!$B$49,"Supportive",IF(C34>Controls!$B$50,"Negative","Neutral")))',None),
     (36,'Sizing flex applied to section C (Supportive only)','=IF(C35="Supportive",Controls!$B$51,0)','0%'),
     (37,'Largest-impact variable (position / new-renewal / industry / ISO)','=IF(MAX($I$22:$I$25)=0,"none rated",INDEX($B$22:$B$25,MATCH(MAX($I$22:$I$25),$I$22:$I$25,0))&": "&TEXT(INDEX($F$22:$F$25,MATCH(MAX($I$22:$I$25),$I$22:$I$25,0)),"0.0%")&" PLR vs "&TEXT($C$10,"0.0%")&" book")',None)]
for r,t,f,fmt in res: label(sc,f'B{r}',t); put(sc,f'C{r}',f,B,fmt,align=LEFT)
sc['C35'].font=Font(name='Calibri',size=12,bold=True)
note(sc,'B39','Largest-impact: column I = |PLR - book PLR| for the four single-dimension rows; the biggest gap wins. Display only.')
sc.sheet_view.showGridLines=False; sc.freeze_panes='C14'

# ======================= Comparable Deals =======================
cd=wb.create_sheet('Comparable Deals')
cd['A1']='COMPARABLE DEALS  -  the whole book with two flags; filter column V to 1 to see the primary cohort'; cd['A1'].font=Font(name='Calibri',size=14,bold=True)
cd['A2']=f"=\"Primary cohort: \"&{HS}$C$30&\"   |   \"&{HS}$C$31&\" seasoned deals\""; cd['A2'].font=B
ch=['Contract ID','DBA','Status','Funded','Position','Industry','ISO','New/renewal','Factor','Advance','Collected','% paid','Net cash','Principal lost','Term (bd)','Age (bd)','Days on book','Seasoned','Default','Open','Best tier (1 = closest)','In primary cohort']
sub(cd,4,ch); cd.row_dimensions[4].height=30
src=['A','B','D','E','G','Y','K','L','M','N','Q','S','AJ','AK','AA','AB','AN','AC','AE','AH','BB','BC']
for i in range(2,NR+1):
    r=3+i
    for j,col in enumerate(src):
        c=cd.cell(r,1+j,f"={HDR}{col}{i}"); c.font=N
        if col in ('N','Q','AJ','AK'): c.number_format='"$"#,##0'
        if col=='E': c.number_format='mm/dd/yyyy'
        if col=='S': c.number_format='0%'
        if col=='M': c.number_format='0.00'
cd.auto_filter.ref=f'A4:V{NR+3}'
widths(cd,{'A':12,'B':28,'C':18,'D':11,'E':8,'F':26,'G':24,'H':11,'I':7,'J':11,'K':11,'L':7,'M':11,'N':11,'O':8,'P':8,'Q':9,'R':9,'S':8,'T':7,'U':10,'V':10})
cd.freeze_panes='C5'

# ======================= README =======================
rd=wb.create_sheet('README',0); widths(rd,{'A':3,'B':125})
txt=["ASPIRE COMBINED UNDERWRITING MODEL  v4"+("  -  DEMO: 8 sample deals, Deal sheet populated" if DEMO else ""),"",
"HOW TO RUN A DEAL","Deal sheet: sections A, B and existing positions as always; DEAL PROFILE (right of the bank data): ISO, industry, new/renewal, credit score. Position comes from C42.",
"Read C (cash flow), C2 (house-book indicators, same colour code), D (decision on the requested offer), E (max offer at the sizing line). Section F is an outline of the final offer, to be built.",
"","HOW THE HOUSE BOOK CHANGES THE OFFER","1. Historical Score matches the deal to seven cohorts (new/renewal + industry + position ... down to the whole seasoned book) and picks the LOWEST tier with enough seasoned deals and dollars (Controls section 4).",
"2. Index = that cohort's principal lost rate / the seasoned book's. Outcome: SUPPORTIVE (index <= 0.85), NEUTRAL, NEGATIVE (index > 1.15). Thresholds on Controls section 5.",
"3. Supportive moves the sizing lines (column H of section C) part-way toward the red lines, so section E sizes a larger max and metrics may show YELLOW, never RED. Neutral = the tool exactly as built. Negative = green lines, then a fund cut in section F (planned).",
"","SIDE NOTE","A blend rate and a credibility factor (small cohorts pulled toward the book) are under consideration; not built.",
"","COLOUR KEY (colour-blind safe)","Editable = bold black on grey with a thick border. Italic grey = link to Controls. Statuses are always written as words (GREEN / YELLOW / RED / KNOCKOUT) as well as coloured."]
for i,t in enumerate(txt):
    c=rd.cell(2+i,2,t); c.font=B if (t.isupper() or i==0) else N; c.alignment=Alignment(wrap_text=True,vertical='top')
rd.sheet_view.showGridLines=False

# ---------- demo population ----------
if DEMO:
    d=ws
    d['B6']=40000; d['C6']=1.45; d['D6']=24; d['E6']='Weekly'
    for i,(m,rev,bal,neg,nsf) in enumerate([('Aug',88000,7400,0,0),('Jul',84000,6900,0,1),('Jun',91000,8100,0,0),('May',82000,6500,1,0)]):
        r=12+i; d[f'B{r}']=m; d[f'C{r}']=rev; d[f'D{r}']=bal; d[f'E{r}']=neg; d[f'F{r}']=nsf
    d['K12']='Fundwell'; d['K13']='Eating & Drinking Places'; d['K14']='New Deal'; d['K15']=672

for shn in wb.sheetnames:
    sh=wb[shn]; sh.sheet_properties.pageSetUpPr.fitToPage=True; sh.page_setup.fitToWidth=1; sh.page_setup.fitToHeight=0
wb._sheets=[wb[n] for n in ['README','Deal','Historical Score','Comparable Deals','Historical Data','Controls','Ref']]
wb.active=1; wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print('saved',OUT,'rows',NR-1)
