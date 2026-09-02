"""Fill test inputs into the model, recalc with LibreOffice, print the outputs."""
import openpyxl, subprocess, json, sys, os
SK='/root/.claude/skills/synced/eb90e4dd-cd3e-43fc-a79c-53774f60d4a9_b639b35e-cd02-43f4-b1b3-b2389b9963f6/xlsx/scripts/recalc.py'
SRC=sys.argv[1]; OUTDIR=sys.argv[2]; os.makedirs(OUTDIR,exist_ok=True)
ONLY=sys.argv[3].split(',') if len(sys.argv)>3 else None
def months(ws,rows):
    for i,(m,rev,bal,neg,nsf) in enumerate(rows):
        r=12+i; ws[f'B{r}']=m; ws[f'C{r}']=rev; ws[f'D{r}']=bal; ws[f'E{r}']=neg; ws[f'F{r}']=nsf
    for r in range(12+len(rows),24):
        ws[f'B{r}']=None; ws[f'C{r}']=None; ws[f'D{r}']=None; ws[f'E{r}']=None; ws[f'F{r}']=None
def positions(ws,rows):
    for r in range(28,34):
        ws[f'B{r}']=None; ws[f'C{r}']=None; ws[f'D{r}']=None; ws[f'G{r}']=None
    for i,(f,fund,pmt,freq) in enumerate(rows):
        r=28+i; ws[f'B{r}']=f; ws[f'C{r}']=fund; ws[f'D{r}']=pmt; ws[f'E{r}']=freq
SCEN={
 'A_new3rd_construction_USC8':dict(offer=(75000,1.49,30,'Weekly'),months=[('Jul',124000,9200,0,0),('Jun',118000,8100,1,0),('May',121000,9800,0,1),('Apr',115000,8900,0,0)],pos=[('Fox Funding',60000,2500,'Weekly')],prof=dict(K12='United Secured Capital 8',K13='General Building Contractors',K14='TX',K15='New Deal',K16=640)),
 'B_renewal_restaurant_Westwood':dict(offer=(60000,1.45,26,'Weekly'),months=[('Jul',150000,22000,0,0),('Jun',142000,19500,0,0),('May',147000,21000,0,0),('Apr',139000,18000,0,0)],pos=[('Aspire (prior)',50000,2100,'Weekly')],prof=dict(K12='Westwood',K13='Eating & Drinking Places',K14='TX',K15='Renewal',K16=705)),
 'C_knockout_nsf':dict(offer=(40000,1.49,100,'Daily'),months=[('Jul',60000,1500,5,3),('Jun',58000,1200,4,2),('May',61000,1700,2,1)],pos=[],prof=dict(K12='Fundwell',K13='Trucking & Warehousing',K14='FL',K15='New Deal',K16=610)),
 'D_zero_revenue':dict(offer=(None,None,None,None),months=[],pos=[],prof=dict(K12='Fundwell',K13='Health Services',K14='CA',K15='New Deal',K16=680)),
 'E_unknown_ISO_no_industry':dict(offer=(30000,1.45,20,'Weekly'),months=[('Jul',80000,6000,0,0),('Jun',82000,6500,0,0),('May',79000,6100,0,0)],pos=[('OnDeck LOC',None,1800,'Monthly')],prof=dict(K12='Brand New ISO LLC',K13=None,K14='NY',K15='New Deal',K16=690)),
 'F_small_cohort_Triumph_pos5_renewal':dict(offer=(25000,1.49,80,'Daily'),months=[('Jul',70000,5000,0,0),('Jun',71000,5200,1,0),('May',69000,4800,0,0)],pos=[('A',None,150,'Daily'),('B',None,120,'Daily'),('C',None,200,'Daily'),('D',None,90,'Daily')],prof=dict(K12='Triumph Cap Fund',K13='Personal Services',K14='GA',K15='Renewal',K16=660)),
 'G_daily_health_PMF':dict(offer=(50000,1.49,120,'Daily'),months=[('Aug',210000,14000,0,0),('Jul',205000,15500,0,0),('Jun',198000,13800,0,0),('May',215000,16000,0,0),('Apr',190000,12000,0,0),('Mar',188000,12500,0,0)],pos=[('Kapitus',139300,14782,'Monthly'),('Mulligan',191561,6777.8,'Weekly')],prof=dict(K12='PMF',K13='Health Services',K14='OH',K15='New Deal',K16=745)),
 'H_six_positions':dict(offer=(20000,1.49,60,'Daily'),months=[('Jul',90000,4000,0,0),('Jun',88000,4200,0,0)],pos=[('A',None,300,'Daily'),('B',None,250,'Daily'),('C',None,200,'Daily'),('D',None,150,'Daily'),('E',None,100,'Daily')],prof=dict(K12='Ballicci',K13='Miscellaneous Retail',K14='NJ',K15='New Deal',K16=650)),
 'I_credit_red_and_trend_red':dict(offer=(40000,1.49,100,'Daily'),months=[('Jul',70000,5000,0,0),('Jun',85000,6000,0,0),('May',100000,7000,0,0)],pos=[('X',None,900,'Weekly')],prof=dict(K12='Fundwell',K13='Auto Repair & Services',K14='CA',K15='New Deal',K16=590)),
 'J_position_override_strong':dict(offer=(100000,1.40,36,'Weekly'),months=[('Jul',400000,90000,0,0),('Jun',410000,95000,0,0),('May',390000,88000,0,0),('Apr',405000,92000,0,0)],pos=[],prof=dict(K12='Premium Capital (AMRK)',K13='Legal Services',K14='IL',K15='New Deal',K16=760,K18=3)),
}
for name,sc in SCEN.items():
    if ONLY and name not in ONLY: continue
    wb=openpyxl.load_workbook(SRC); ws=wb['Deal']
    b,c,d,e=sc['offer']; ws['B6']=b; ws['C6']=c; ws['D6']=d; ws['E6']=e
    months(ws,sc['months']); positions(ws,sc['pos'])
    for k,v in sc['prof'].items(): ws[k]=v
    out=f'{OUTDIR}/{name}.xlsx'; wb.save(out)
    r=subprocess.run(['python3',SK,out,'300'],capture_output=True,text=True); j=json.loads(r.stdout)
    v=openpyxl.load_workbook(out,data_only=True); D=v['Deal']; S=v['Decision Summary']; H=v['Historical Score']; O=v['Offer Engine']; K=v['Checks']
    print('\n'+'='*20,name,'| recalc:',j.get('status'),j.get('total_errors'), j.get('error_summary',''))
    print(' L1:',D['C52'].value,'|',D['D2'].value)
    print(' statuses:',[D[f'F{r}'].value for r in range(38,46)],'score',D['C55'].value,'/',D['D55'].value,'scaled',round(D['D57'].value or 0,2),D['C56'].value)
    print(' CF max daily',D['C60'].value,'|',D['D60'].value,'| term',D['C62'].value,'factor',D['D62'].value,'| CF max fund',D['C63'].value)
    print(' primary tier',H['C29'].value,H['C30'].value,'| n',H['C31'].value,'| lost',H['C33'].value,'| Z',H['C37'].value,'| blended',H['C38'].value,'| idx',H['C39'].value,'| lvl',H['C40'].value,'->',H['C42'].value,H['C43'].value,'| conf',H['C41'].value)
    for r in range(14,27):
        if H[f'E{r}'].value: print('   tier',H[f'A{r}'].value,'n',H[f'E{r}'].value,'/',H[f'D{r}'].value,'lost',H[f'L{r}'].value,'Z',round(H[f'V{r}'].value or 0,2),'idx',H[f'X{r}'].value,'seasoned-enough',H[f'U{r}'].value)
    print(' OE adjusted: fund',O['D8'].value,'factor',O['D7'].value,'term',O['D6'].value,'daily',O['D9'].value,'freq',O['D16'].value,'| statuses',[O[f'E{r}'].value for r in range(26,34)],'| proofs',[O[f'D{r}'].value for r in range(36,41)])
    print(' FINAL:',S['C5'].value,'|',S['D5'].value)
    for r in (6,7,8,9,10,11,16,18,19,21,23,24,26,28): print('   ',S[f'B{r}'].value,':',S[f'C{r}'].value,'|',S[f'D{r}'].value)
    print(' CONCLUSION:',S['C30'].value)
    print(' Checks:',K['B2'].value)
    fails=[(K[f'B{r}'].value,K[f'D{r}'].value) for r in range(4,60) if K[f'C{r}'].value=='FAIL']; print(' FAILS:',fails)
