"""v2: history can loosen sizing (flex toward yellow), C2 indicator block, no stips / confidence / conclusion.
Usage: build_model_v2.py OUT.xlsx [demo]   demo = 8 sample deals + populated Deal sheet + demo minimums"""
import openpyxl, datetime as dt, copy, sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter as L
from openpyxl.cell.cell import MergedCell
from openpyxl.workbook.properties import CalcProperties
ROOT='/home/user/Underwriting-Tool-House-Book'
UW=f'{ROOT}/originals/Updated_Underwriting_Tool (v2 upload).xlsx'
FULL=f'{ROOT}/originals/House Book Analytics (2).xlsx'
SAMPLE=f'{ROOT}/originals/Copy_of_House_Book_Analytics (sample).xlsx'
OUT=sys.argv[1]; DEMO=len(sys.argv)>2 and sys.argv[2]=='demo'

DARK=PatternFill('solid',fgColor='1F2937'); MID=PatternFill('solid',fgColor='4B5563'); GREY=PatternFill('solid',fgColor='E5E7EB'); INPUT=PatternFill('solid',fgColor='FFF9C4')
WHITE=Font(name='Calibri',size=11,bold=True,color='FFFFFF'); SUB=Font(name='Calibri',size=9,bold=True,color='FFFFFF')
B=Font(name='Calibri',size=11,bold=True); N=Font(name='Calibri',size=11); NOTE=Font(name='Calibri',size=9,color='374151'); LINK=Font(name='Calibri',size=11,color='008000'); INP=Font(name='Calibri',size=11,color='0000FF')
thin=Side(style='thin',color='9CA3AF'); med=Side(style='medium',color='374151')
BOX=Border(left=thin,right=thin,top=thin,bottom=thin); IBOX=Border(left=med,right=med,top=med,bottom=med)
CEN=Alignment(horizontal='center',vertical='center'); LEFT=Alignment(horizontal='left',vertical='center'); WRAP=Alignment(horizontal='left',vertical='top',wrap_text=True)
G_FILL=PatternFill('solid',fgColor='00B050'); Y_FILL=PatternFill('solid',fgColor='FFFF00'); R_FILL=PatternFill('solid',fgColor='FF0000')
def rag_rules(ws,rng):
    for v,f,fc in (('GREEN',G_FILL,'FFFFFF'),('YELLOW',Y_FILL,'000000'),('RED',R_FILL,'FFFFFF'),('KNOCKOUT',R_FILL,'FFFFFF')):
        ws.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=[f'"{v}"'],fill=f,font=Font(color=fc,bold=True)))
def dec_rules(ws,rng):
    for v,f,fc in (('APPROVE',G_FILL,'FFFFFF'),('CONDITIONAL',Y_FILL,'000000'),('DECLINE',R_FILL,'FFFFFF'),('MANUAL REVIEW',PatternFill('solid',fgColor='F97316'),'FFFFFF')):
        ws.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=[f'"{v}"'],fill=f,font=Font(color=fc,bold=True)))
def chk_rules(ws,rng):
    for v,f,fc in (('OK',G_FILL,'FFFFFF'),('WARN',Y_FILL,'000000'),('FAIL',R_FILL,'FFFFFF')):
        ws.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=[f'"{v}"'],fill=f,font=Font(color=fc,bold=True)))
def hdr(ws,cell,text,span=None):
    ws[cell]=text; ws[cell].font=WHITE; ws[cell].fill=DARK; ws[cell].alignment=LEFT
    if span: ws.merge_cells(span)
def sub(ws,row,labels,start=1):
    for i,t in enumerate(labels):
        c=ws.cell(row,start+i,t); c.font=SUB; c.fill=MID; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=BOX
def put(ws,addr,val,font=N,fmt=None,align=CEN,border=BOX,fill=None):
    c=ws[addr]; c.value=val; c.font=font; c.alignment=align; c.border=border
    if fmt: c.number_format=fmt
    if fill: c.fill=fill
    return c
def label(ws,addr,text,font=N):
    c=ws[addr]; c.value=text; c.font=font; c.alignment=LEFT; c.border=BOX; return c
def note(ws,addr,text):
    c=ws[addr]; c.value=text; c.font=NOTE; c.alignment=LEFT; return c
def widths(ws,d):
    for k,v in d.items(): ws.column_dimensions[k].width=v

# ---------- house-book data ----------
hv=openpyxl.load_workbook(FULL,data_only=True); ref=hv['Ref']
sic=[(ref.cell(r,1).value,ref.cell(r,2).value,ref.cell(r,3).value) for r in range(4,68)]
hols=[ref.cell(r,6).value for r in range(4,45) if ref.cell(r,6).value]
fullraw=[list(r) for r in hv['Data'].iter_rows(min_row=2,max_row=1068,max_col=27,values_only=True) if r[0] is not None]
isos=sorted({r[12] for r in fullraw if r[12]},key=str.lower); states=sorted({r[11] for r in fullraw if r[11]})
statuses=[('Closed','No','Yes','Yes','No'),('Open','No','No','No','Yes'),('Legal','Yes','No','No','No'),('In House Collection','Yes','No','No','No'),('Lowered Payments','No','No','No','No'),('Written-Off','Yes','Yes','No','No'),('Bankruptcy','Yes','Yes','No','No'),('Settlement Agreement','Yes','Yes','No','No'),('Closed/NSF','No','Yes','Yes','No'),('Closed/NSF Unpaid','No','Yes','Yes','No'),('At-Risk','Yes','No','No','No')]
industries=sorted({s[1] for s in sic})+['No SIC','Unclassified']
raw=[list(r) for r in openpyxl.load_workbook(SAMPLE,data_only=True)['Data'].iter_rows(min_row=2,max_row=9,max_col=27,values_only=True)] if DEMO else fullraw
NR=1+len(raw)  # last data row
HDR="'Historical Data'!"; HS="'Historical Score'!"; OE="'Offer Engine'!"; DS="'Decision Summary'!"; C='Controls!'
def hr(col): return f"{HDR}${col}$2:${col}${NR}"

# ---------- Deal sheet from the template ----------
wb=openpyxl.load_workbook(UW)
for n in wb.sheetnames:
    if n!='New Template': del wb[n]
ws=wb['New Template']; ws.title='Deal'
def shift(ws,first,last,n):
    merges=[m for m in list(ws.merged_cells.ranges) if first<=m.min_row<=last]
    for m in merges: ws.unmerge_cells(str(m))
    for r in range(last,first-1,-1):
        for c in range(2,14):
            ws.cell(r+n,c)._style=copy.copy(ws.cell(r,c)._style); ws.cell(r+n,c).value=None
        if ws.row_dimensions[r].height: ws.row_dimensions[r+n].height=ws.row_dimensions[r].height
    for m in merges: ws.merge_cells(start_row=m.min_row+n,start_column=m.min_col,end_row=m.max_row+n,end_column=m.max_col)
    for r in range(first,first+n):
        for c in range(2,14):
            if not isinstance(ws.cell(r,c),MergedCell): ws.cell(r,c).value=None
shift(ws,45,63,1)          # credit-score row
for c in range(2,14): ws.cell(45,c)._style=copy.copy(ws.cell(44,c)._style)
shift(ws,49,64,10)         # C2 house-book indicator block
for r in range(49,59):     # style C2 rows like the metric rows
    for c in range(2,14): ws.cell(r,c)._style=copy.copy(ws.cell(38 if r>=51 else (36 if r==49 else 37),c)._style)
for m in list(ws.merged_cells.ranges):
    if m.min_row in (11,18) and m.min_col>=10: ws.unmerge_cells(str(m))
for r in range(11,19):
    for c in range(10,14):
        if not isinstance(ws.cell(r,c),MergedCell): ws.cell(r,c).value=None
for r in range(59,75):
    for c in range(2,14):
        if not isinstance(ws.cell(r,c),MergedCell): ws.cell(r,c).value=None
ws['E35']=None
ws.conditional_formatting=ConditionalFormattingList()
ws.data_validations.dataValidation=[dv for dv in ws.data_validations.dataValidation if str(dv.sqref) in ('E28:E33','B12:B23','E6')]

# rows (after shifts): metrics 38-45, info 46-47, C2 49-57, D 59-62, E 64-73, F 75+
ws['C2']='=C62'
ws['D2']='=IF(C62="","",C60&" red flag(s) vs limit "&C61&"   |   "&COUNTIF(F38:F45,"GREEN")&"G / "&COUNTIF(F38:F45,"YELLOW")&"Y / "&COUNTIF(F38:F45,"RED")&"R   |   score "&TEXT(C65,"0.0")&"/"&D65&" = "&C66&"   |   FINAL: "&IF(C77="","-",C77))'
ws['F6']=f'={C}$B$15'; ws['F6'].font=Font(name='Calibri',size=12,bold=True,color='008000')
ws['J6']='=IF($E$6="Weekly","Weekly remittance","Daily remittance")'
ws['K6']='=IF(OR($E$6="",C8="",C24=0),"",IF($E$6="Weekly",D8,C8))'
ws['L6']='=IF(OR(L7="",C24=0),"",L7*C24/IF($E$6="Weekly",$L$8,$L$9))'
ws['L7']='=IF(OR($E$6="",C8="",C24=0),"",ROUND(IF($E$6="Weekly",D8*$L$8,C8*$L$9)/C24*100,2)/100)'
ws['K8']=f'={C}$B$28'; ws['K9']=f'={C}$B$29'; ws['L8']=f'={C}$B$30'; ws['L9']=f'={C}$B$31'
for a in ['K8','K9','L8','L9']: ws[a].font=LINK
ws['E8']='=IF(C8="","",C8*$K$9)'; ws['F8']='=IF(OR(D6="",D6=0),"",IF(E6="Weekly",D6,D6/5)/($K$9/5))'
# Deal profile J11:M16
hdr(ws,'J11','DEAL PROFILE  (house-book match inputs)','J11:M11')
prof=[('ISO','K12','dropdown'),('Industry (SIC major group)','K13','dropdown'),('State  (optional - informational only)','K14','optional'),('New deal or renewal','K15','REQUIRED for a cohort'),('Credit score (FICO)','K16','number')]
for i,(t,a,n_) in enumerate(prof):
    r=12+i; c=ws.cell(r,10,t); c.font=Font(name='Calibri',size=9,bold=True,color='FFFFFF'); c.fill=MID; c.alignment=LEFT; c.border=BOX
    cell=ws[a]; cell.border=IBOX; cell.alignment=CEN; cell.font=Font(name='Calibri',size=11,bold=True); cell.fill=INPUT
ws['L12']=f'=IF(K12="","",COUNTIF({hr("K")},K12)&" deals in book")'
ws['L13']=f'=IF(K13="","",COUNTIF({hr("Y")},K13)&" deals in book")'
ws['L14']=f'=IF(K14="","not needed for an offer",COUNTIF({hr("J")},K14)&" deals in book")'
ws['L15']='=IF(K15="","required for history match","")'
ws['L16']=f'=IF(K16="","not scored",IF(K16>={C}$B$12,"GREEN",IF(K16>={C}$C$12,"YELLOW","RED")))'
for r in range(12,17): ws.merge_cells(f'L{r}:M{r}'); ws[f'L{r}'].font=NOTE
note(ws,'J17','Position for matching = C42 (existing positions + 1). Bands and thresholds: Controls.')
for a,src in [('K12','=Ref!$E$2:$E$'+str(1+len(isos))),('K13','=Ref!$G$2:$G$'+str(1+len(industries))),('K14','=Ref!$I$2:$I$'+str(1+len(states))),('K15','"New Deal,Renewal"')]:
    dv=DataValidation(type='list',formula1=src,allow_blank=True); dv.add(a); ws.add_data_validation(dv)
for r in range(28,34):
    ws[f'F{r}']=f'=IF(D{r}<>"",IF(E{r}="Daily",D{r},IF(E{r}="Weekly",D{r}/5,D{r}/$K$9)),IF(C{r}="","",C{r}*{C}$B$33/({C}$B$23*$K$9)))'
    ws[f'H{r}']=f'=IF(OR(C{r}="",D{r}=""),"",ROUND(C{r}*{C}$B$33/D{r},0))'
    ws[f'I{r}']=f'=IF(H{r}="","",H{r}/IF(E{r}="Daily",$K$9,IF(E{r}="Weekly",$K$9/5,1)))'
    ws[f'J{r}']=f'=IF(OR(I{r}="",G{r}=""),"",MAX(0,I{r}-(TODAY()-G{r})/{C}$B$32))'
ws['C34']='=SUMPRODUCT(((C28:C33<>"")+(D28:D33<>"")>0)*(E28:E33<>"Monthly"))'
ws['F34']='=E34*$K$9'
# C. metrics
ws['H37']='Sizing line'; ws['H37']._style=copy.copy(ws['G37']._style)
metrics=[(38,'Total LVG (all debt / rev)',5,'=IF(C24=0,"",(E34+IF(C8="",0,C8))*$K$9/C24)','0.00%'),(39,'New holdback SP%',6,'=IF(OR(C24=0,E8=""),"",E8/C24)','0.00%'),(40,'Daily debt / balance',7,'=IF(D24=0,"",(E34+IF(C8="",0,C8))/D24)','0.00%'),(41,'Fund / revenue',8,'=IF(OR(C24=0,B6=""),"",B6/C24)','0.00'),(42,'# positions (incl. new)',9,'=IF(C24=0,"",C34+1)','0'),(43,'Revenue trend (avg MoM)',10,'=G24','0.0%'),(44,'Neg days + NSFs (total)',11,'=E24+F24','0'),(45,'Credit score (FICO)',12,'=IF(K16="","",K16)','0')]
for r,t,cr,f,fmt in metrics:
    ws[f'B{r}']=t; ws[f'C{r}']=f; ws[f'C{r}'].number_format=fmt
    for col,cc in (('D','B'),('E','C'),('G','D')):
        ws[f'{col}{r}']=f'={C}${cc}${cr}'; ws[f'{col}{r}'].font=LINK; ws[f'{col}{r}'].number_format='0.0%' if r in (38,39,40,43) else ('0.00' if r==41 else '0')
    if r==45: ws[f'F{r}']=f'=IF(C{r}="","",IF(AND(G{r}<>0,C{r}<=G{r}),"KNOCKOUT",IF(C{r}>=D{r},"GREEN",IF(C{r}<E{r},"RED","YELLOW"))))'
    elif r==43: ws['F43']='=IF(C24=0,"",IF(NOT(ISNUMBER(C43)),"YELLOW",IF(C43<=G43,"KNOCKOUT",IF(C43>=D43,"GREEN",IF(C43<=E43,"RED","YELLOW")))))'
    else: ws[f'F{r}']=f'=IF(C{r}="","",IF(C{r}>=G{r},"KNOCKOUT",IF(C{r}<=D{r},"GREEN",IF(C{r}>=E{r},"RED","YELLOW"))))'
    if r<=41:  # sizing line = green + flex x (red - green)
        ws[f'H{r}']=f'=D{r}+{HS}$C$42*(E{r}-D{r})'; ws[f'H{r}']._style=copy.copy(ws[f'C{r}']._style); ws[f'H{r}'].number_format=ws[f'D{r}'].number_format; ws[f'H{r}'].font=Font(name='Calibri',size=11,bold=True,color='3730A3')
ws['B46']='Existing leverage (before new)'; ws['C46']='=IF(C24=0,"",F34/C24)'; ws['D46']='-'; ws['E46']='-'; ws['F46']='INFO'
ws['B47']='Balance / revenue'; ws['C47']='=IF(C24=0,"",D24/C24)'; ws['D47']='-'; ws['E47']='-'; ws['F47']='INFO'
ws['C46'].number_format='0.0%'; ws['C47'].number_format='0.0%'
# C2. house-book indicators
ws['B49']='C2.  HOUSE-BOOK INDICATORS   (set the sizing line + modifier; NOT counted in the cash-flow score)'
for h,t in zip('BCDEFG',['Cohort','Blended lost rate','Green <=','Red >','Status','Seasoned n']): ws[f'{h}50']=t
book=f'{HS}$C$10'
def ind(r,lbl,hsrow):
    ws[f'B{r}']=lbl; ws[f'C{r}']=f'=IF({HS}$L${hsrow}="","",{HS}$W${hsrow})'; ws[f'C{r}'].number_format='0.0%'
    ws[f'D{r}']=f'={book}*{C}$A$50'; ws[f'E{r}']=f'={book}*{C}$A$51'; ws[f'D{r}'].number_format=ws[f'E{r}'].number_format='0.0%'; ws[f'D{r}'].font=ws[f'E{r}'].font=LINK
    ws[f'F{r}']=f'=IF({HS}$E${hsrow}=0,"NO DEALS",IF({HS}$L${hsrow}="","TOO FEW",IF({HS}$X${hsrow}<={C}$A$50,"GREEN",IF({HS}$X${hsrow}<={C}$A$51,"YELLOW","RED"))))'
    ws[f'G{r}']=f'={HS}$E${hsrow}'; ws[f'G{r}'].number_format='0'
ws['B51']=f'="Primary cohort: "&{HS}$C$30'; ws['C51']=f'=IF({HS}$C$29=7,"",{HS}$C$38)'; ws['C51'].number_format='0.0%'
ws['D51']=f'={book}*{C}$A$50'; ws['E51']=f'={book}*{C}$A$51'; ws['D51'].number_format=ws['E51'].number_format='0.0%'; ws['D51'].font=ws['E51'].font=LINK
ws['F51']=f'=IF({HS}$C$29=7,"NO EVIDENCE",IF({HS}$C$39<={C}$A$50,"GREEN",IF({HS}$C$39<={C}$A$51,"YELLOW","RED")))'; ws['G51']=f'={HS}$C$31'
ind(52,'ISO',25); ind(53,'Industry',24); ind(54,'New deal / renewal',23); ind(55,'Position',22); ind(56,'State  (informational - never sizes the offer)',26)
ws['B57']='Seasoned deals behind the primary cohort'; ws['C57']=f'=IF({HS}$C$29=7,0,INDEX({HS}$E$14:$E$20,{HS}$C$29))'; ws['C57'].number_format='0'
ws['D57']=f'={C}$B$44'; ws['E57']=f'={C}$B$40'; ws['D57'].font=ws['E57'].font=LINK; ws['D57'].number_format=ws['E57'].number_format='0'
ws['F57']='=IF(C57>=D57,"GREEN",IF(C57>=E57,"YELLOW","RED"))'
ws['G57']=f'="modifier: "&{HS}$C$41&"  |  sizing flex "&TEXT({HS}$C$42,"0%")'; ws['G57'].alignment=LEFT
ws.merge_cells('G57:H57')
rag_rules(ws,'F38:F45'); rag_rules(ws,'F51:F57'); dec_rules(ws,'C2'); dec_rules(ws,'C62'); dec_rules(ws,'C77')
# D. decision
ws['B59']='D.  DECISION  (cash-flow layer, on the requested offer)'
ws['B60']='Red flags'; ws['C60']='=COUNTIF(F38:F45,"RED")+COUNTIF(F38:F45,"KNOCKOUT")'; ws['D60']='Knockouts'; ws['E60']='=COUNTIF(F38:F45,"KNOCKOUT")'
ws['B61']='Decline if red flags >=  (Controls)'; ws['C61']=f'={C}$B$13'; ws['D61']='Minimum fund $'; ws['E61']=f'={C}$B$14'
ws['C61'].font=Font(name='Calibri',size=12,bold=True,color='008000'); ws['E61'].font=Font(name='Calibri',size=12,bold=True,color='008000')
ws['B62']='DECISION'
ws['C62']='=IF(C24=0,"",IF(OR(E60>0,IF(ISNUMBER(C73),C73,0)<E61),"DECLINE",IF(OR(B6="",B6=0),"",IF(C60>=C61,"DECLINE",IF(COUNTIF(F38:F45,"GREEN")=D65,"APPROVE","CONDITIONAL")))))'
ws['D62']='=IF(C62="DECLINE",IF(E60>0,"KNOCKOUT: "&INDEX(B38:B45,MATCH("KNOCKOUT",F38:F45,0))&" is past its knockout line.",IF(C60>=C61,"Red flags reached the threshold.","Max fund "&IF(ISNUMBER(C73),TEXT(C73,"$#,##0"),"$0")&" is below your minimum "&TEXT(E61,"$#,##0")&".")),IF(C62="APPROVE","All "&D65&" decision metrics green.",IF(C62="","Enter offer + bank data.","Fundable but not clean - review yellows/reds above.")))'
# E. sizing at the SIZING LINE (= green line when the modifier is neutral)
ws['B64']='E.  DEAL QUALITY  &  MAXIMUM OFFER at the sizing line  (green line unless the house book is supportive)'
ws['B65']='Score (G=1, Y=0.5, R=0)'; ws['C65']='=COUNTIF(F38:F45,"GREEN")+0.5*COUNTIF(F38:F45,"YELLOW")'; ws['D65']='=SUMPRODUCT(--(F38:F45<>""))'; ws['E65']='scored metrics'
ws['B66']='Tier  (score scaled to 7 for the bands)'; ws['C66']=f'=IF(C24=0,"",IF(D67>={C}$B$25,"Strong",IF(D67>={C}$C$25,"Solid",IF(D67>={C}$D$25,"Moderate","Weak"))))'
ws['B67']='Max daily by SP% sizing line'; ws['C67']='=IF(C24=0,"",H39*C24/$K$9)'; ws['D67']='=IF(D65=0,0,C65*7/D65)'; ws['E67']='scaled score /7'
ws['B68']='Max daily by Total LVG sizing line'; ws['C68']='=IF(C24=0,"",(H38*C24-F34)/$K$9)'
ws['B69']='Max daily by Daily/Balance sizing line'; ws['C69']='=IF(OR(C24=0,D24=0),"",H40*D24-E34)'
ws['B70']='MAX DAILY  (weakest link)'; ws['C70']='=IF(C24=0,"",MAX(0,MIN(C67:C69)))'
ws['D70']='=IF(C70="","",IF(C70=0,"no clean room - existing debt is over the sizing line",INDEX({"SP% holdback";"Total LVG";"Daily/Balance"},MATCH(C70,C67:C69,0))&" is the limit"))'
ws['B71']='Weekly / monthly equivalent'; ws['C71']='=IF(OR(C70="",C70=0),"",C70*5)'; ws['D71']='=IF(OR(C70="",C70=0),"",C70*$K$9)'
ws['B72']='Term (daily pmts) / factor from bands'
ws['C72']=f'=IF(OR(C70="",C70=0,C65=""),"",MAX(20,MIN(INDEX({C}$B$19:$B$22,MATCH(D67,{C}$A$19:$A$22,1)),FLOOR({C}$B$23*$K$9,5),FLOOR(H41*C24*D72/C70,5))))'
ws['D72']=f'=IF(OR(C65="",C24=0),"",INDEX({C}$C$19:$C$22,MATCH(D67,{C}$A$19:$A$22,1)))'
ws['B73']='MAX FUND at the sizing line'
ws['C73']='=IF(COUNTIF(F42:F45,"KNOCKOUT")>0,"-",IF(OR(C70="",C70=0,C72="",D72=""),"",FLOOR(C70*C72/D72,500)))'
ws['D73']='=IF(C73="-","KNOCKOUT on the merchant - do not fund.",IF(C73="","","payback "&TEXT(C73*D72,"$#,##0")&" over "&C72&" daily pmts ("&TEXT(C72/5,"0")&" wks)  |  green-line max "&IF(ISNUMBER('+"'Offer Engine'"+'!$C$8),TEXT('+"'Offer Engine'"+'!$C$8,"$#,##0"),"-")))'
for a in ['C67','C68','C69','C70']: ws[a].number_format='"$"#,##0.00'
ws['D67'].number_format='0.0'; ws['C71'].number_format=ws['D71'].number_format='"$"#,##0'; ws['C72'].number_format='0'; ws['D72'].number_format='0.00'; ws['C73'].number_format='"$"#,##0'
# F. final offer
r0=75
hdr(ws,f'B{r0}','F.  FINAL OFFER  (max at the sizing line, then the modifier)   -   full detail on Decision Summary',f'B{r0}:H{r0}')
sub(ws,r0+1,['Item','Value','Detail'],start=2); ws.merge_cells(f'D{r0+1}:H{r0+1}')
rows=[('FINAL DECISION',5),('Recommended fund',6),('Factor',7),('Term',8),('Payment',9),('Frequency',10),('Sizing line used',11),('Holdback % at final offer',12),('Historical modifier',17),('Comparable deals (seasoned / all)',19),('Cohort principal lost rate',21),('Manual-review triggers',25)]
for i,(t,dsr) in enumerate(rows):
    r=r0+2+i
    c=ws.cell(r,2,t); c.font=B; c.alignment=Alignment(horizontal='right',vertical='center'); c.border=BOX
    c=ws.cell(r,3,f'={DS}$C${dsr}'); c.font=Font(name='Calibri',size=12 if i<2 else 11,bold=True); c.alignment=CEN; c.border=BOX
    c=ws.cell(r,4,f'={DS}$D${dsr}'); c.font=NOTE; c.alignment=LEFT; ws.merge_cells(f'D{r}:H{r}')
ws['C78'].number_format='"$"#,##0'; ws['C79'].number_format='0.00'; ws['C81'].number_format='"$"#,##0.00'; ws['C84'].number_format='0.0%'; ws['C87'].number_format='0.0%'
note(ws,f'B{r0+2+len(rows)}','Layer 1 stays the ceiling: the sizing line never reaches the red line, knockouts are never removed, the final fund never exceeds the max at the sizing line. See Checks.')
ws.column_dimensions['J'].width=27; ws.column_dimensions['K'].width=26; ws.column_dimensions['L'].width=16; ws.column_dimensions['M'].width=12; ws.column_dimensions['H'].width=13
ws.sheet_view.showGridLines=False

# ======================= Controls =======================
cs=wb.create_sheet('Controls'); widths(cs,{'A':52,'B':15,'C':15,'D':15,'E':15,'F':13,'G':13,'H':13,'I':60})
cs['A1']='CONTROLS  -  every editable threshold, tier, historical weight and rule'; cs['A1'].font=Font(name='Calibri',size=14,bold=True)
note(cs,'A2','Blue = editable. Do not insert or delete rows; formulas reference fixed cells.')
hdr(cs,'A3','1. CASH-FLOW THRESHOLDS','A3:I3'); sub(cs,4,['Metric','Green (at or better)','Red (at or worse)','Knockout','Direction','','','','Note'])
thr=[('Total LVG (all debt / rev)',0.30,0.45,0.55,'higher is worse','0.0%',''),('New holdback SP%',0.18,0.25,0.35,'higher is worse','0.0%',''),('Daily debt / balance',0.18,0.30,0.35,'higher is worse','0.0%',''),('Fund / revenue',0.60,0.90,1.25,'higher is worse','0.00',''),('# positions (incl. new)',2,5,6,'higher is worse','0',''),('Revenue trend (avg MoM)',-0.05,-0.15,-0.30,'lower is worse','0.0%',''),('Neg days + NSFs (total)',1,6,10,'higher is worse','0','total across all months entered'),('Credit score (FICO)',650,630,0,'lower is worse','0','650+ green, 630-649 yellow, below 630 red; knockout 0 = off')]
for i,(t,g,r_,k,d,fmt,n_) in enumerate(thr):
    r=5+i; label(cs,f'A{r}',t); put(cs,f'B{r}',g,INP,fmt,border=IBOX); put(cs,f'C{r}',r_,INP,fmt,border=IBOX); put(cs,f'D{r}',k,INP,fmt,border=IBOX); put(cs,f'E{r}',d,N,align=LEFT); note(cs,f'I{r}',n_)
label(cs,'A13','Decline if red flags (RED + KNOCKOUT) >='); put(cs,'B13',2,INP,'0',border=IBOX)
label(cs,'A14','Minimum fund $'); put(cs,'B14',10000,INP,'"$"#,##0',border=IBOX)
label(cs,'A15','Underwriting fee %'); put(cs,'B15',0.05,INP,'0%',border=IBOX)
hdr(cs,'A17','2. PRICING BANDS','A17:I17'); sub(cs,18,['Score >=','Term cap (daily pmts)','Factor','','','','','','Note'])
for i,(s,t,f) in enumerate([(0,100,1.49),(5,130,1.45),(6,150,1.45),(7,195,1.40)]):
    r=19+i; put(cs,f'A{r}',s,INP,'0.0',border=IBOX); put(cs,f'B{r}',t,INP,'0',border=IBOX); put(cs,f'C{r}',f,INP,'0.00',border=IBOX)
label(cs,'A23','Max term (months)'); put(cs,'B23',9,INP,'0',border=IBOX)
label(cs,'A24','Approved factor tiers (low -> high)'); put(cs,'B24',1.40,INP,'0.00',border=IBOX); put(cs,'C24',1.45,INP,'0.00',border=IBOX); put(cs,'D24',1.49,INP,'0.00',border=IBOX)
label(cs,'A25','Tier cut-offs on 7-point scale: Strong / Solid / Moderate'); put(cs,'B25',6.5,INP,'0.0',border=IBOX); put(cs,'C25',5,INP,'0.0',border=IBOX); put(cs,'D25',3.5,INP,'0.0',border=IBOX)
hdr(cs,'A27','3. CONVERSION CONSTANTS','A27:I27')
for r,(t,v,fmt) in enumerate([('Weeks per month (ours)',4.331,'0.000'),('Business days per month (ours)',21.655,'0.000'),('MoneyBadger weeks per month',4,'0'),('MoneyBadger business days per month',21,'0'),('Calendar days per month',30.4375,'0.0000'),('Assumed factor on existing positions',1.45,'0.00')],start=28):
    label(cs,f'A{r}',t); put(cs,f'B{r}',v,INP,fmt,border=IBOX)
hdr(cs,'A35','4. HISTORICAL EVIDENCE SETTINGS','A35:I35')
hist=[('As-of date for seasoning (date of the CRM export)',dt.datetime(2026,8,25),'mm/dd/yyyy',''),('Window: funded on or after',dt.datetime(2024,11,1),'mm/dd/yyyy',''),('Window: funded on or before',dt.datetime(2026,8,25),'mm/dd/yyyy','12/31/2025 = settled book'),
      ('Seasoning multiple (age >= multiple x contracted term)',1.25,'0.00',''),('Minimum seasoned deals for a cohort to be PRIMARY',3 if DEMO else 10,'0','RECOMMENDED for the full book: 10' if DEMO else ''),('Minimum seasoned funded $ for a cohort to be PRIMARY',30000 if DEMO else 150000,'"$"#,##0','RECOMMENDED for the full book: $150,000' if DEMO else ''),
      ('Credibility constant K (seasoned deals for half weight)',30,'0','Z = n / (n + K); small cohorts are pulled toward the book'),('Minimum seasoned deals before a rate is shown',2 if DEMO else 5,'0','RECOMMENDED for the full book: 5' if DEMO else ''),('Seasoned deals for the sample-size indicator to be GREEN',5 if DEMO else 30,'0','RECOMMENDED for the full book: 30' if DEMO else '')]
for i,(t,v,fmt,n_) in enumerate(hist):
    r=36+i; label(cs,f'A{r}',t); put(cs,f'B{r}',v,INP,fmt,border=IBOX); note(cs,f'I{r}',n_)
hdr(cs,'A47','5. HISTORICAL MODIFIER TABLE   index = blended cohort lost rate / seasoned-book lost rate  (1.00 = book)','A47:I47')
sub(cs,48,['Index at or below','Level','Sizing flex (0 = green line, 1 = red line)','Fund multiplier','Term step (bands)','Factor step (tiers)','Frequency','Manual review?','Note'])
mods=[(0.60,'Strong positive',0.90,1.00,0,0,'As proposed','N','RECOMMENDED: size 90% of the way to the red line - metrics may go yellow, never red'),
      (0.85,'Mild positive',0.45,1.00,0,0,'As proposed','N','half-way to the red line'),
      (1.15,'Neutral',0.00,1.00,0,0,'As proposed','N','equals the cash-flow tool exactly as built'),
      (1.40,'Mild negative',0.00,0.85,-1,1,'As proposed','N',''),
      (1.75,'Negative',0.00,0.70,-1,1,'Weekly','N',''),
      (999,'Severe',0.00,0.55,-2,1,'Weekly','Y','sent to manual review')]
for i,m in enumerate(mods):
    r=49+i
    for j,v in enumerate(m[:8]):
        put(cs,f'{L(1+j)}{r}',v,B if j==1 else INP,{0:'0.00',2:'0%',3:'0.00',4:'0',5:'0'}.get(j),border=BOX if j==1 else IBOX)
    note(cs,f'I{r}',m[8])
note(cs,'A55','The C2 block on the Deal sheet uses A50 (0.85) and A51 (1.15) as its GREEN / RED lines for every house-book indicator. Modifier = Neutral whenever no cohort meets the minimums.')
hdr(cs,'A57','6. STATUS MAPPING','A57:I57'); sub(cs,58,['Status','Default?','Resolved?','Closed?','Open?','Deals in file','','',''])
for i,s in enumerate(statuses):
    r=59+i; label(cs,f'A{r}',s[0])
    for j in range(4): put(cs,f'{L(2+j)}{r}',s[j+1],INP,border=IBOX)
    put(cs,f'F{r}',f'=COUNTIF({hr("D")},A{r})',N,'0')
hdr(cs,'A72','7. MANUAL-REVIEW TRIGGERS  (Y = active)','A72:I72'); sub(cs,73,['Trigger','Active?','Parameter','','','','','','Note'])
trig=[('Historical modifier is Severe','Y',None,''),('Credit score RED AND another metric RED','Y',None,''),('Positions incl. new at or above','Y',4,'fund up to a 4th'),('New deal, position at/above parameter, construction industry','Y',3,''),('ISO or industry has no deals in the house book','Y',None,''),('Final fund below minimum fund $','Y',None,''),('No cohort met the minimums (book baseline in use)','N',None,'off by default')]
for i,(t,a,p,n_) in enumerate(trig):
    r=74+i; label(cs,f'A{r}',t); put(cs,f'B{r}',a,INP,border=IBOX)
    if p is not None: put(cs,f'C{r}',p,INP,'0',border=IBOX)
    note(cs,f'I{r}',n_)
cs.sheet_view.showGridLines=False

# ======================= Ref =======================
rf=wb.create_sheet('Ref'); widths(rf,{'A':8,'B':40,'C':32,'E':44,'G':40,'I':10,'K':14,'M':22})
for col,t in zip('ABCEGIKM',['SIC2','Industry (SIC major group)','Division','ISO (house-book spelling)','Industry list','State','Fed holiday','Status list']):
    c=rf[f'{col}1']; c.value=t; c.font=SUB; c.fill=MID; c.alignment=CEN; c.border=BOX
for i,(a,b,c_) in enumerate(sic): rf.cell(2+i,1,a); rf.cell(2+i,2,b); rf.cell(2+i,3,c_)
for i,v in enumerate(isos): rf.cell(2+i,5,v)
for i,v in enumerate(industries): rf.cell(2+i,7,v)
for i,v in enumerate(states): rf.cell(2+i,9,v)
for i,v in enumerate(hols): rf.cell(2+i,11,v).number_format='mm/dd/yyyy'
for i,s in enumerate(statuses): rf.cell(2+i,13,s[0])
HOL=f'Ref!$K$2:$K${1+len(hols)}'

# ======================= Historical Data =======================
hs=wb.create_sheet('Historical Data')
cols=['Contract ID','DBA','Legal name','Status','Start date','Last active','Position','SIC','SIC2','State','ISO','New or renewal','Factor','Advance','Commission','Payback','Collected','Balance','% paid','Exp dur (raw)','NSF count','Avg rev/mo','Holdback %','Biz start',
      'Industry','Pay frequency','Term (bus. days)','Age at as-of (bus. days)','Seasoned','In window','Default','Resolved','Closed family','Open','Status valid','Net cash','Principal lost','Adv / revenue','Has revenue','Days on book (bus.)','Renewal',
      'Eligible (seasoned & in window)','m: position','m: new/renewal','m: industry','m: ISO','m: state','T1 NR+Ind+Pos','T2 NR+ISO+Pos','T3 NR+Ind','T4 NR+ISO','T5 NR+Pos','T6 NR','T7 Book','Best tier','In primary cohort','Seq']
for i,t in enumerate(cols):
    c=hs.cell(1,1+i,t); c.font=SUB; c.fill=MID if i<24 else DARK; c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=BOX
hs.row_dimensions[1].height=42
keep=[0,1,2,3,4,5,6,7,9,11,12,13,14,15,16,17,18,19,20,21,23,24,25,26]
st="Controls!$A$59:$E$69"
for i,r in enumerate(raw):
    rr=2+i
    for j,k in enumerate(keep):
        v=r[k]; v=str(v) if (k in (7,9) and v is not None) else v
        hs.cell(rr,1+j,v)
    f={'Y':f'=IF(I{rr}="","No SIC",IFERROR(VLOOKUP(I{rr},Ref!$A$2:$B$65,2,FALSE),"Unclassified"))','Z':f'=IF(T{rr}="","",IF(T{rr}<12,"Daily","Weekly"))',
       'AA':f'=IF(T{rr}="",0,IF(T{rr}<12,T{rr}*Controls!$B$29,T{rr}*5))','AB':f'=IF(E{rr}="",0,NETWORKDAYS(E{rr},Controls!$B$36,{HOL}))','AC':f'=IF(AND(AA{rr}>0,AB{rr}>=Controls!$B$39*AA{rr}),1,0)',
       'AD':f'=IF(AND(E{rr}<>"",E{rr}>=Controls!$B$37,E{rr}<=Controls!$B$38),1,0)','AE':f'=IF(IFERROR(VLOOKUP(D{rr},{st},2,FALSE),"No")="Yes",1,0)','AF':f'=IF(IFERROR(VLOOKUP(D{rr},{st},3,FALSE),"No")="Yes",1,0)',
       'AG':f'=IF(IFERROR(VLOOKUP(D{rr},{st},4,FALSE),"No")="Yes",1,0)','AH':f'=IF(IFERROR(VLOOKUP(D{rr},{st},5,FALSE),"No")="Yes",1,0)','AI':f'=IF(ISNUMBER(MATCH(D{rr},Controls!$A$59:$A$69,0)),1,0)',
       'AJ':f'=Q{rr}-N{rr}','AK':f'=IF(AE{rr}=1,MAX(0,N{rr}-Q{rr}),0)','AL':f'=IF(OR(V{rr}="",V{rr}=0),0,N{rr}/V{rr})','AM':f'=IF(OR(V{rr}="",V{rr}=0),0,1)','AN':f'=IF(OR(E{rr}="",F{rr}=""),0,NETWORKDAYS(E{rr},F{rr},{HOL}))','AO':f'=IF(L{rr}="Renewal",1,0)','AP':f'=AC{rr}*AD{rr}',
       'AQ':f"=IF(AND(G{rr}<>\"\",{HS}$C$6<>\"\"),IF(G{rr}={HS}$C$6,1,0),0)",'AR':f"=IF(AND(L{rr}<>\"\",{HS}$C$5<>\"\"),IF(L{rr}={HS}$C$5,1,0),0)",'AS':f"=IF(AND(Y{rr}<>\"\",{HS}$C$7<>\"\"),IF(Y{rr}={HS}$C$7,1,0),0)",'AT':f"=IF(AND(K{rr}<>\"\",{HS}$C$8<>\"\"),IF(K{rr}={HS}$C$8,1,0),0)",'AU':f"=IF(AND(J{rr}<>\"\",{HS}$C$9<>\"\"),IF(J{rr}={HS}$C$9,1,0),0)",
       'AV':f'=AR{rr}*AS{rr}*AQ{rr}','AW':f'=AR{rr}*AT{rr}*AQ{rr}','AX':f'=AR{rr}*AS{rr}','AY':f'=AR{rr}*AT{rr}','AZ':f'=AR{rr}*AQ{rr}','BA':f'=AR{rr}','BB':'=1',
       'BC':f'=IF(AV{rr}=1,1,IF(AW{rr}=1,2,IF(AX{rr}=1,3,IF(AY{rr}=1,4,IF(AZ{rr}=1,5,IF(BA{rr}=1,6,7))))))','BD':f"=IF(AD{rr}=1,CHOOSE({HS}$C$29,AV{rr},AW{rr},AX{rr},AY{rr},AZ{rr},BA{rr},BB{rr}),0)",'BE':f'=IF(BD{rr}=1,COUNTIF($BD$2:BD{rr},1),"")'}
    for col,fx in f.items(): hs[f'{col}{rr}']=fx
    for col in ('E','F','X'): hs[f'{col}{rr}'].number_format='mm/dd/yyyy'
    for col in ('N','O','P','Q','R'): hs[f'{col}{rr}'].number_format='"$"#,##0'
    hs[f'S{rr}'].number_format='0.0%'; hs[f'W{rr}'].number_format='0.0%'
hs.freeze_panes='E2'; widths(hs,{'A':12,'B':26,'C':26,'D':18,'K':24,'Y':26})

# ======================= Historical Score =======================
sc=wb.create_sheet('Historical Score'); widths(sc,{'A':7,'B':36,'C':34})
for col in range(4,28): sc.column_dimensions[L(col)].width=10
sc['A1']='HISTORICAL SCORE  -  comparable cohorts, credibility and the modifier'; sc['A1'].font=Font(name='Calibri',size=14,bold=True)
note(sc,'A2','Rates use SEASONED deals only (age >= multiple x contracted term). Blended = Z x cohort + (1-Z) x seasoned book, Z = n/(n+K). Index = blended / book.')
hdr(sc,'A4','DEAL BEING MATCHED','A4:C4')
for r,(t,f) in enumerate([('New deal or renewal','=IF(Deal!$K$15="","",Deal!$K$15)'),('Position (Deal C42)','=IF(Deal!$C$42="","",Deal!$C$42)'),('Industry','=IF(Deal!$K$13="","",Deal!$K$13)'),('ISO','=IF(Deal!$K$12="","",Deal!$K$12)'),('State (informational)','=IF(Deal!$K$14="","",Deal!$K$14)')],start=5):
    label(sc,f'B{r}',t); put(sc,f'C{r}',f,B)
label(sc,'B10','Seasoned-book principal lost rate (baseline)'); put(sc,'C10','=IF($L$20="",0,$L$20)',B,'0.0%')
label(sc,'B11','Seasoned-book deals / funded'); put(sc,'C11','=TEXT($E$20,"#,##0")&" / "&TEXT($G$20,"$#,##0")',B)
heads=['Tier','Cohort','Keys','Deals (all)','Seasoned','% seasoned','Funded $ (seasoned)','Open n','Open %','Default n','Default %','Principal lost rate','Collected on defaults','Return on funded','Avg adv/rev','Avg term (bd)','Avg factor','Avg deal size','Avg days on book','Renewal share','Meets minimums?','Credibility Z','Blended lost rate','Relative index','Meets min (1/0)','Impact vs book','Rate shown?']
sub(sc,13,heads); sc.row_dimensions[13].height=40
def cohort_row(r,tier,name,flag,keys):
    F=hr(flag); E_=hr('AP')
    put(sc,f'A{r}',tier,B); label(sc,f'B{r}',name); label(sc,f'C{r}',keys)
    fx={'D':f'=SUMPRODUCT({F}*{hr("AD")})','E':f'=SUMPRODUCT({F}*{E_})','F':f'=IF(D{r}=0,"",E{r}/D{r})','G':f'=SUMPRODUCT({F}*{E_}*{hr("N")})','H':f'=SUMPRODUCT({F}*{E_}*{hr("AH")})','I':f'=IF(E{r}=0,"",H{r}/E{r})','J':f'=SUMPRODUCT({F}*{E_}*{hr("AE")})','K':f'=IF(E{r}=0,"",J{r}/E{r})',
        'L':f'=IF(OR(E{r}<Controls!$B$43,G{r}=0),"",SUMPRODUCT({F}*{E_}*{hr("AK")})/G{r})','M':f'=IF(OR(E{r}<Controls!$B$43,SUMPRODUCT({F}*{E_}*{hr("AE")}*{hr("N")})=0),"",SUMPRODUCT({F}*{E_}*{hr("AE")}*{hr("Q")})/SUMPRODUCT({F}*{E_}*{hr("AE")}*{hr("N")}))',
        'N':f'=IF(OR(E{r}<Controls!$B$43,G{r}=0),"",SUMPRODUCT({F}*{E_}*{hr("AJ")})/G{r})','O':f'=IF(SUMPRODUCT({F}*{E_}*{hr("AM")})=0,"",SUMPRODUCT({F}*{E_}*{hr("AM")}*{hr("AL")})/SUMPRODUCT({F}*{E_}*{hr("AM")}))',
        'P':f'=IF(E{r}=0,"",SUMPRODUCT({F}*{E_}*{hr("AA")})/E{r})','Q':f'=IF(E{r}=0,"",SUMPRODUCT({F}*{E_}*{hr("M")})/E{r})','R':f'=IF(E{r}=0,"",G{r}/E{r})','S':f'=IF(E{r}=0,"",SUMPRODUCT({F}*{E_}*{hr("AN")})/E{r})','T':f'=IF(E{r}=0,"",SUMPRODUCT({F}*{E_}*{hr("AO")})/E{r})',
        'U':f'=IF(AND(E{r}>=Controls!$B$40,G{r}>=Controls!$B$41),"Yes","No")','V':f'=IF(E{r}+Controls!$B$42=0,0,E{r}/(E{r}+Controls!$B$42))','W':f'=IF(L{r}="",$C$10,V{r}*L{r}+(1-V{r})*$C$10)','X':f'=IF(OR($C$10="",$C$10=0),"",W{r}/$C$10)',
        'Y':f'=IF(U{r}="Yes",1,0)','Z':f'=IF(L{r}="",0,ABS(W{r}-$C$10))','AA':f'=IF(L{r}="","No","Yes")'}
    for col,f in fx.items(): put(sc,f'{col}{r}',f,N)
    for col in 'FIKLMNTW': sc[f'{col}{r}'].number_format='0.0%'
    for col in 'GR': sc[f'{col}{r}'].number_format='"$"#,##0'
    for col in 'OQVX': sc[f'{col}{r}'].number_format='0.00'
    sc[f'Z{r}'].number_format='0.0%'; sc[f'P{r}'].number_format='0'; sc[f'S{r}'].number_format='0'
tiers=[(1,'T1: new/renewal + industry + position','AV','NR + industry + position'),(2,'T2: new/renewal + ISO + position','AW','NR + ISO + position'),(3,'T3: new/renewal + industry','AX','NR + industry'),(4,'T4: new/renewal + ISO','AY','NR + ISO'),(5,'T5: new/renewal + position','AZ','NR + position'),(6,'T6: new/renewal only','BA','NR'),(7,'T7: seasoned book (baseline)','BB','all seasoned deals in window')]
for t,nme,flag,keys in tiers: cohort_row(13+t,t,nme,flag,keys)
sub(sc,21,['','SINGLE-DIMENSION EVIDENCE  (C2 block on Deal; largest-impact test)']+['']*25)
for i,(nme,flag) in enumerate([('Position only','AQ'),('New/renewal only','AR'),('Industry only','AS'),('ISO only','AT'),('State only (INFORMATIONAL)','AU')]): cohort_row(22+i,'S'+str(i+1),nme,flag,nme)
sc.conditional_formatting.add('U14:U26',CellIsRule(operator='equal',formula=['"Yes"'],fill=G_FILL,font=Font(color='FFFFFF',bold=True)))
sc.conditional_formatting.add('U14:U26',CellIsRule(operator='equal',formula=['"No"'],fill=Y_FILL))
hdr(sc,'A28','RESULT','A28:C28')
res=[(29,'Primary tier (first tier meeting the minimums)','=IFERROR(MATCH(1,$Y$14:$Y$20,0),7)','0'),(30,'Primary cohort','=INDEX($B$14:$B$20,$C$29)',None),(31,'Seasoned deals / all matches','=TEXT(INDEX($E$14:$E$20,$C$29),"#,##0")&" / "&TEXT(INDEX($D$14:$D$20,$C$29),"#,##0")',None),
     (32,'% of comparable deals seasoned','=INDEX($F$14:$F$20,$C$29)','0.0%'),(33,'Cohort principal lost rate (seasoned)','=INDEX($L$14:$L$20,$C$29)','0.0%'),(34,'Cohort open share','=INDEX($I$14:$I$20,$C$29)','0.0%'),(35,'Cohort return on funded capital','=INDEX($N$14:$N$20,$C$29)','0.0%'),(36,'Cohort collected on defaults','=INDEX($M$14:$M$20,$C$29)','0.0%'),
     (37,'Credibility Z','=INDEX($V$14:$V$20,$C$29)','0.00'),(38,'Blended lost rate','=INDEX($W$14:$W$20,$C$29)','0.0%'),(39,'Relative index (blended / book)','=INDEX($X$14:$X$20,$C$29)','0.00'),
     (40,'Modifier level (1 strong positive .. 3 neutral .. 6 severe)','=IF(OR(C29=7,C39=""),3,COUNTIF(Controls!$A$49:$A$54,"<"&C39)+1)','0'),(41,'Modifier level name','=INDEX(Controls!$B$49:$B$54,C40)',None),
     (42,'Sizing flex (0 = green line, 1 = red line)','=INDEX(Controls!$C$49:$C$54,C40)','0%'),(43,'Fund multiplier','=INDEX(Controls!$D$49:$D$54,C40)','0.00'),(44,'Term step (bands)','=INDEX(Controls!$E$49:$E$54,C40)','0'),(45,'Factor step (tiers)','=INDEX(Controls!$F$49:$F$54,C40)','0'),(46,'Frequency rule','=INDEX(Controls!$G$49:$G$54,C40)',None),(47,'Manual review flag from modifier','=INDEX(Controls!$H$49:$H$54,C40)',None),
     (48,'Largest-impact variable (position / NR / industry / ISO)','=IF(MAX($Z$22:$Z$25)=0,"none rated",INDEX($B$22:$B$25,MATCH(MAX($Z$22:$Z$25),$Z$22:$Z$25,0)))',None),
     (49,'   its blended lost rate vs book','=IF(MAX($Z$22:$Z$25)=0,"",TEXT(INDEX($W$22:$W$25,MATCH(MAX($Z$22:$Z$25),$Z$22:$Z$25,0)),"0.0%")&" vs "&TEXT($C$10,"0.0%")&" book ("&TEXT(INDEX($E$22:$E$25,MATCH(MAX($Z$22:$Z$25),$Z$22:$Z$25,0)),"0")&" seasoned deals)")',None),
     (50,'ISO','=IF($C$8="","not entered",IF($E$25=0,$C$8&": no seasoned deals in the book",IF($L$25="",$C$8&": "&$E$25&" seasoned deals - too few to rate",$C$8&": "&$E$25&" seasoned deals, lost "&TEXT($L$25,"0.0%")&" vs "&TEXT($C$10,"0.0%")&" book, Z "&TEXT($V$25,"0.00")&" -> "&IF($X$25<=Controls!$A$50,"stronger than book",IF($X$25<=Controls!$A$51,"neutral","weaker than book")))))',None),
     (51,'Industry','=IF($C$7="","not entered",IF($E$24=0,$C$7&": no seasoned deals in the book",IF($L$24="",$C$7&": "&$E$24&" seasoned deals - too few to rate",$C$7&": "&$E$24&" seasoned deals, lost "&TEXT($L$24,"0.0%")&" vs "&TEXT($C$10,"0.0%")&" book, Z "&TEXT($V$24,"0.00")&" -> "&IF($X$24<=Controls!$A$50,"stronger than book",IF($X$24<=Controls!$A$51,"neutral","weaker than book")))))',None),
     (52,'New deal / renewal','=IF($C$5="","not entered",IF($L$23="",$C$5&": "&$E$23&" seasoned deals - too few to rate",$C$5&": "&$E$23&" seasoned deals, lost "&TEXT($L$23,"0.0%")&" vs "&TEXT($C$10,"0.0%")&" book -> "&IF($X$23<=Controls!$A$50,"stronger than book",IF($X$23<=Controls!$A$51,"neutral","weaker than book"))))',None),
     (53,'Position','=IF($C$6="","not entered",IF($L$22="","position "&$C$6&": "&$E$22&" seasoned deals - too few to rate","position "&$C$6&": "&$E$22&" seasoned deals, lost "&TEXT($L$22,"0.0%")&" vs "&TEXT($C$10,"0.0%")&" book -> "&IF($X$22<=Controls!$A$50,"stronger than book",IF($X$22<=Controls!$A$51,"neutral","weaker than book"))))',None),
     (54,'State (informational only)','=IF($C$9="","not entered - not needed",IF($E$26=0,$C$9&": no seasoned deals",IF($L$26="",$C$9&": "&$E$26&" seasoned deals - too few to rate",$C$9&": "&$E$26&" seasoned deals, lost "&TEXT($L$26,"0.0%")&" vs "&TEXT($C$10,"0.0%")&" book - informational only")))',None),
     (55,'Conflicting signals? (ISO vs industry on opposite sides)','=IF(OR($L$24="",$L$25=""),"No",IF(OR(AND($X$24<Controls!$A$50,$X$25>Controls!$A$51),AND($X$25<Controls!$A$50,$X$24>Controls!$A$51)),"YES - review both cohorts","No"))',None),
     (56,'Warnings','=IF($C$29=7,"No cohort met the minimums - book baseline, modifier Neutral; ","")&IF(C29<>1,"Tighter tiers 1-"&(C29-1)&" too small; ","")&IF(C55<>"No","Conflicting ISO / industry signals; ","")',None)]
for r,t,f,fmt in res: label(sc,f'B{r}',t); put(sc,f'C{r}',f,B,fmt,align=LEFT)
sc.sheet_view.showGridLines=False; sc.freeze_panes='D14'

# ======================= Offer Engine =======================
oe=wb.create_sheet('Offer Engine'); widths(oe,{'A':3,'B':44,'C':18,'D':18,'E':18,'F':14,'G':56})
oe['B1']='OFFER ENGINE  -  green-line base  ->  sized at the house-book sizing line  ->  final after the modifier'; oe['B1'].font=Font(name='Calibri',size=14,bold=True)
note(oe,'B2','Base = the cash-flow tool exactly as built (green lines). Sized = Deal section E (sizing line, moved toward the red line only when the book is supportive). Final = sized x multiplier, term / factor steps. Final daily payment never exceeds the sized max daily.')
sub(oe,4,['','Item','Base (green line)','Sized (sizing line)','Final (modifier)','Status','How'],start=1)
D='Deal!'
K9=f'{D}$K$9'
rows=[(5,'Max daily payment',f'=IF({D}$C$24=0,"",MAX(0,MIN({D}$D$39*{D}$C$24/{K9},({D}$D$38*{D}$C$24-{D}$F$34)/{K9},IF({D}$D$24=0,0,{D}$D$40*{D}$D$24-{D}$E$34))))',f'={D}$C$70',f'={D}$C$70','"$"#,##0.00','weakest of holdback / total LVG / daily-debt-to-balance at that line'),
      (6,'Term (daily payments)',f'=IF(OR(C5="",C5=0),"",MAX(20,MIN(INDEX(Controls!$B$19:$B$22,MATCH({D}$D$67,Controls!$A$19:$A$22,1)),FLOOR(Controls!$B$23*{K9},5),FLOOR({D}$D$41*{D}$C$24*C7/C5,5))))',f'={D}$C$72','=IF(OR(D6="",D8="-"),"",MAX(20,MIN(D6,INDEX(Controls!$B$19:$B$22,MAX(1,MATCH(Deal!$D$67,Controls!$A$19:$A$22,1)+C11)))))','0','band term; final steps down the band table'),
      (7,'Factor',f'={D}$D$72',f'={D}$D$72','=IF(D7="","",INDEX(Controls!$B$24:$D$24,MIN(3,MAX(1,MATCH(D7,Controls!$B$24:$D$24,0)+C12))))','0.00','next approved tier when stepped; never above 1.49'),
      (8,'Fund $',f'=IF(COUNTIF({D}$F$42:$F$45,"KNOCKOUT")>0,"-",IF(OR(C5="",C5=0,C6="",C7=""),"",FLOOR(C5*C6/C7,500)))',f'={D}$C$73','=IF(D8="-","-",IF(OR(D8="",E6="",E7=""),"",MAX(0,FLOOR(MIN(D8*C10,D5*E6/E7),500))))','"$"#,##0','final = MIN(sized x multiplier, sized max daily x final term / final factor)'),
      (9,'Daily payment','=IF(OR(C8="",C8="-"),"",C8*C7/C6)','=IF(OR(D8="",D8="-"),"",D8*D7/D6)','=IF(OR(E8="",E8="-",E8=0),"",E8*E7/E6)','"$"#,##0.00','fund x factor / term'),
      (10,'Fund multiplier (Controls)',f'={HS}$C$43',None,None,'0.00',''),(11,'Term step',f'={HS}$C$44',None,None,'0',''),(12,'Factor step',f'={HS}$C$45',None,None,'0',''),(13,'Frequency rule',f'={HS}$C$46',None,None,None,''),(14,'Manual review flag',f'={HS}$C$47',None,None,None,''),(15,'Sizing flex',f'={HS}$C$42',None,None,'0%','share of the way from the green line to the red line used for sizing'),
      (16,'Frequency (final)',f'=IF({D}$E$6="","Daily",{D}$E$6)',f'=IF({D}$E$6="","Daily",{D}$E$6)',f'=IF(C13="Weekly","Weekly",IF({D}$E$6="","Daily",{D}$E$6))',None,''),
      (17,'Term in final units','=IF(OR(C6="",C8="-"),"",IF(C16="Weekly",C6/5,C6))','=IF(OR(D6="",D8="-"),"",IF(D16="Weekly",D6/5,D6))','=IF(OR(E6="",E8="-"),"",IF(E16="Weekly",E6/5,E6))','0','weeks or daily payments'),
      (18,'Payment in final units','=IF(C9="","",IF(C16="Weekly",C9*5,C9))','=IF(D9="","",IF(D16="Weekly",D9*5,D9))','=IF(E9="","",IF(E16="Weekly",E9*5,E9))','"$"#,##0.00',''),
      (19,'Monthly payment','=IF(C9="","",C9*Controls!$B$29)','=IF(D9="","",D9*Controls!$B$29)','=IF(E9="","",E9*Controls!$B$29)','"$"#,##0',''),
      (20,'Payback','=IF(OR(C8="",C8="-"),"",C8*C7)','=IF(OR(D8="",D8="-"),"",D8*D7)','=IF(OR(E8="",E8="-"),"",E8*E7)','"$"#,##0',''),
      (21,'Net disbursed','=IF(OR(C8="",C8="-"),"",C8*(1-Controls!$B$15))','=IF(OR(D8="",D8="-"),"",D8*(1-Controls!$B$15))','=IF(OR(E8="",E8="-"),"",E8*(1-Controls!$B$15))','"$"#,##0',''),
      (22,'Term (months)','=IF(C6="","",C6/Controls!$B$29)','=IF(D6="","",D6/Controls!$B$29)','=IF(E6="","",E6/Controls!$B$29)','0.0','')]
for r,t,cb,cs_,ce,fmt,how in rows:
    label(oe,f'B{r}',t); put(oe,f'C{r}',cb,N,fmt)
    if cs_ is not None: put(oe,f'D{r}',cs_,N,fmt)
    if ce is not None: put(oe,f'E{r}',ce,B,fmt)
    note(oe,f'G{r}',how)
hdr(oe,'B24','METRICS AT THE FINAL OFFER  (yellow allowed when the book is supportive; RED / KNOCKOUT never)','B24:G24')
sub(oe,25,['','Metric','Base','Sized','Final','Status (final)','Green / Red / Knockout'],start=1)
def m3(r,t,expr,cr,fmt):
    label(oe,f'B{r}',t)
    for col,src in (('C','C'),('D','D'),('E','E')): put(oe,f'{col}{r}',expr.replace('#',src),B if col=='E' else N,fmt)
    put(oe,f'F{r}',f'=IF(E{r}="","",IF(E{r}>=Controls!$D${cr},"KNOCKOUT",IF(E{r}<=Controls!$B${cr},"GREEN",IF(E{r}>=Controls!$C${cr},"RED","YELLOW"))))',B)
    put(oe,f'G{r}',f'=TEXT(Controls!$B${cr},"{fmt}")&" / "&TEXT(Controls!$C${cr},"{fmt}")&" / "&TEXT(Controls!$D${cr},"{fmt}")',NOTE,align=LEFT)
m3(26,'Total LVG (all debt / rev)',f'=IF(OR({D}$C$24=0,#9=""),"",({D}$E$34+#9)*Controls!$B$29/{D}$C$24)',5,'0.0%')
m3(27,'New holdback SP%',f'=IF(OR({D}$C$24=0,#9=""),"",#9*Controls!$B$29/{D}$C$24)',6,'0.0%')
m3(28,'Daily debt / balance',f'=IF(OR({D}$D$24=0,#9=""),"",({D}$E$34+#9)/{D}$D$24)',7,'0.0%')
m3(29,'Fund / revenue',f'=IF(OR({D}$C$24=0,#8="",#8="-"),"",#8/{D}$C$24)',8,'0.00')
for r,t,cr in [(30,'# positions (incl. new)',42),(31,'Revenue trend (avg MoM)',43),(32,'Neg days + NSFs (total)',44),(33,'Credit score (FICO)',45)]:
    label(oe,f'B{r}',t); put(oe,f'E{r}',f'={D}$C${cr}',N); put(oe,f'F{r}',f'={D}$F${cr}',B); note(oe,f'G{r}','merchant-side, unchanged by the offer')
oe['E31'].number_format='0.0%'
rag_rules(oe,'F26:F33')
hdr(oe,'B35','SAFETY PROOF','B35:G35')
for r,t,f in [(36,'Final fund <= sized max fund','=IF(OR(E8="",E8="-",D8="",D8="-"),"OK",IF(E8<=D8,"OK","FAIL"))'),(37,'Final daily payment <= sized max daily','=IF(OR(E9="",D5=""),"OK",IF(E9<=D5+0.005,"OK","FAIL"))'),(38,'No offer-side metric RED / KNOCKOUT at the final offer','=IF(COUNTIF(F26:F29,"RED")+COUNTIF(F26:F29,"KNOCKOUT")=0,"OK","FAIL")'),(39,'Merchant knockouts untouched',f'=IF(COUNTIF({D}$F$42:$F$45,"KNOCKOUT")>0,IF(E8="-","OK","FAIL"),"OK")'),(40,'Sizing flex below 100% (sizing line never reaches the red line)','=IF(C15<1,"OK","FAIL")'),(41,'Red flags at the final offer (offer-side + merchant-side)','=COUNTIF(F26:F33,"RED")+COUNTIF(F26:F33,"KNOCKOUT")')]:
    label(oe,f'B{r}',t); put(oe,f'E{r}',f,B)
chk_rules(oe,'E36:E40'); oe.sheet_view.showGridLines=False

# ======================= Decision Summary =======================
ds=wb.create_sheet('Decision Summary'); widths(ds,{'A':3,'B':46,'C':24,'D':92})
ds['B1']='DECISION SUMMARY'; ds['B1'].font=Font(name='Calibri',size=14,bold=True)
sub(ds,4,['','Item','Value','Detail'],start=1)
S=[(5,'DECISION','=IF(Deal!$C$24=0,"",IF(OR(Deal!$C$73="-",Deal!$C$73="",'+OE+'$E$8="",'+OE+'$E$8=0,'+OE+'$E$8<Controls!$B$14,'+OE+'$E$41>=Controls!$B$13),"DECLINE",IF(C25<>"","MANUAL REVIEW",IF(OR(COUNTIF('+OE+'$F$26:$F$33,"YELLOW")+COUNTIF('+OE+'$F$26:$F$33,"RED")>0,'+HS+'$C$40>=4),"CONDITIONAL","APPROVE"))))',
   '=IF(C5="","Enter bank data + deal profile.",IF(C5="DECLINE",IF(COUNTIF(Deal!$F$42:$F$45,"KNOCKOUT")>0,"Merchant knockout: "&INDEX(Deal!$B$42:$B$45,MATCH("KNOCKOUT",Deal!$F$42:$F$45,0))&".",IF(OR('+OE+'$E$8="",'+OE+'$E$8=0),"No clean room under the sizing line.",IF('+OE+'$E$8<Controls!$B$14,"Final fund "&TEXT('+OE+'$E$8,"$#,##0")&" is below the minimum "&TEXT(Controls!$B$14,"$#,##0")&".","Red flags at the final offer reached the limit."))),IF(C5="MANUAL REVIEW","Triggers: "&C25,IF(C5="APPROVE","All metrics green at the final offer; evidence neutral or better.","Yellows remain at the final offer or the evidence is negative - see the metric rows."))))',None),
   (6,'Recommended funding amount',"='Offer Engine'!$E$8",'=IF(OR(C6="",C6="-"),"","green-line max "&IF(ISNUMBER('+OE+'$C$8),TEXT('+OE+'$C$8,"$#,##0"),"-")&"  |  sized max "&IF(ISNUMBER('+OE+'$D$8),TEXT('+OE+'$D$8,"$#,##0"),"-")&"  |  final = "&TEXT(C6/'+OE+'$D$8,"0%")&" of sized  (requested "&IF(Deal!$B$6="","n/a",TEXT(Deal!$B$6,"$#,##0"))&")")','"$"#,##0'),
   (7,'Recommended factor rate',"='Offer Engine'!$E$7",'=IF(C7="","","band "&TEXT('+OE+'$D$7,"0.00")&IF('+OE+'$E$7>'+OE+'$D$7," -> raised one tier by the modifier",""))','0.00'),
   (8,'Recommended term',"='Offer Engine'!$E$17",'=IF(C8="","",IF('+OE+'$E$16="Weekly","weeks","daily payments")&"  ("&TEXT('+OE+'$E$22,"0.0")&" months; sized "&'+OE+'$D$6&" daily pmts"&IF('+OE+'$E$6<'+OE+'$D$6,", shortened by the modifier","")&")")','0'),
   (9,'Payment',"='Offer Engine'!$E$18",'=IF(C9="","","per "&IF('+OE+'$E$16="Weekly","week","business day")&";  daily-equivalent "&TEXT('+OE+'$E$9,"$#,##0.00")&", monthly "&TEXT('+OE+'$E$19,"$#,##0")&";  payback "&TEXT('+OE+'$E$20,"$#,##0")&", net "&TEXT('+OE+'$E$21,"$#,##0"))','"$"#,##0.00'),
   (10,'Daily or weekly',"='Offer Engine'!$E$16",'=IF('+OE+'$C$13="Weekly","Weekly required by the modifier","As proposed on the Deal sheet")',None),
   (11,'Sizing line used',"='Offer Engine'!$C$15",'=IF(C11=0,"green lines (the cash-flow tool as built)",TEXT(C11,"0%")&" of the way from the green line to the red line - the book is supportive, yellows allowed")','0%'),
   (12,'Holdback %',"='Offer Engine'!$E$27",'=IF(C12="","",'+OE+'$F$27&"  (green "&TEXT(Controls!$B$6,"0%")&" / red "&TEXT(Controls!$C$6,"0%")&")")','0.0%'),
   (13,'Total leverage (all debt / revenue)',"='Offer Engine'!$E$26",'=IF(C13="","",'+OE+'$F$26&"  (green "&TEXT(Controls!$B$5,"0%")&" / red "&TEXT(Controls!$C$5,"0%")&")")','0.0%'),
   (14,'Funding-to-revenue',"='Offer Engine'!$E$29",'=IF(C14="","",'+OE+'$F$29&"  (green "&TEXT(Controls!$B$8,"0.00")&" / red "&TEXT(Controls!$C$8,"0.00")&")")','0.00'),
   (15,'Daily debt-to-balance',"='Offer Engine'!$E$28",'=IF(C15="","",'+OE+'$F$28&"  (green "&TEXT(Controls!$B$7,"0%")&" / red "&TEXT(Controls!$C$7,"0%")&")")','0.0%'),
   (16,'Number of positions (incl. new)','=Deal!$C$42','=IF(C16="","",Deal!$F$42)','0'),
   (17,'Historical modifier',"='Historical Score'!$C$41",'=IF(C17="","","index "&TEXT('+HS+'$C$39,"0.00")&" (1.00 = seasoned book); sizing flex "&TEXT('+HS+'$C$42,"0%")&", fund x "&TEXT('+HS+'$C$43,"0.00")&", term step "&'+HS+'$C$44&", factor step "&'+HS+'$C$45)',None),
   (18,'Cash-flow risk score','=Deal!$C$65','=IF(C18="","","of "&Deal!$D$65&" scored metrics = "&Deal!$C$66&"  (scaled "&TEXT(Deal!$D$67,"0.0")&"/7 for the bands)")','0.0'),
   (19,'Number of comparable deals (seasoned / all)',"='Historical Score'!$C$31","='Historical Score'!$C$30",None),
   (20,'Percentage of comparable deals seasoned',"='Historical Score'!$C$32",'="credibility Z "&TEXT('+HS+'$C$37,"0.00")&" (n / (n + K))"','0.0%'),
   (21,'Comparable-cohort principal lost rate',"='Historical Score'!$C$33",'=IF(C21="","too few seasoned deals to show a rate","vs seasoned book "&TEXT('+HS+'$C$10,"0.0%")&";  blended "&TEXT('+HS+'$C$38,"0.0%")&";  collected on defaults "&IF('+HS+'$C$36="","n/a",TEXT('+HS+'$C$36,"0%")))','0.0%'),
   (22,'Comparable-cohort return on funded capital',"='Historical Score'!$C$35",'="seasoned deals only"','0.0%'),
   (23,'Largest historical impact',"='Historical Score'!$C$48","='Historical Score'!$C$49",None),
   (24,'Metric that limited the offer','=IF(Deal!$D$70="","",Deal!$D$70)','=IF(Deal!$C$70="","",IF(Deal!$C$70=0,"existing debt already exceeds the sizing line",IF(Deal!$C$72=FLOOR(Deal!$H$41*Deal!$C$24*Deal!$D$72/Deal!$C$70,5),"term also capped by the fund/revenue sizing line","term set by the score band")))',None),
   (25,'Manual-review triggers fired','=IF(AND(Controls!$B$74="Y",'+OE+'$C$14="Y"),"Severe historical modifier; ","")&IF(AND(Controls!$B$75="Y",Deal!$F$45="RED",COUNTIF('+OE+'$F$26:$F$32,"RED")>0),"Credit RED plus another RED; ","")&IF(AND(Controls!$B$76="Y",ISNUMBER(Deal!$C$42),IF(ISNUMBER(Deal!$C$42),Deal!$C$42,0)>=Controls!$C$76),"Positions incl. new at/above "&Controls!$C$76&"; ","")&IF(AND(Controls!$B$77="Y",Deal!$K$15="New Deal",ISNUMBER(Deal!$C$42),IF(ISNUMBER(Deal!$C$42),Deal!$C$42,0)>=Controls!$C$77,OR(Deal!$K$13="General Building Contractors",Deal!$K$13="Special Trade Contractors",Deal!$K$13="Heavy Construction")),"New "&Controls!$C$77&"rd+ position construction deal; ","")&IF(AND(Controls!$B$78="Y",OR(AND(Deal!$K$12<>"",COUNTIF('+hr("K")+',Deal!$K$12)=0),AND(Deal!$K$13<>"",COUNTIF('+hr("Y")+',Deal!$K$13)=0))),"ISO or industry not in house book; ","")&IF(AND(Controls!$B$79="Y",ISNUMBER('+OE+'$E$8),IF(ISNUMBER('+OE+'$E$8),'+OE+'$E$8,0)<Controls!$B$14,IF(ISNUMBER('+OE+'$E$8),'+OE+'$E$8,0)>0),"Final fund below minimum; ","")&IF(AND(Controls!$B$80="Y",'+HS+'$C$29=7),"No cohort met the minimums; ","")','="blank = none"',None),
   (26,'Evidence: ISO',"='Historical Score'!$C$50",None,None),(27,'Evidence: industry',"='Historical Score'!$C$51",None,None),(28,'Evidence: new deal / renewal',"='Historical Score'!$C$52",None,None),(29,'Evidence: position',"='Historical Score'!$C$53",None,None),(30,'Evidence: state (informational)',"='Historical Score'!$C$54",None,None),
   (31,'Warnings',"='Historical Score'!$C$56",None,None),
   (32,'Reds at the requested offer','=IF(COUNTIF(Deal!$F$38:$F$45,"RED")+COUNTIF(Deal!$F$38:$F$45,"KNOCKOUT")=0,"none",IF(Deal!$F$38="RED","total leverage; ","")&IF(Deal!$F$39="RED","holdback; ","")&IF(Deal!$F$40="RED","daily debt/balance; ","")&IF(Deal!$F$41="RED","fund/revenue; ","")&IF(Deal!$F$42="RED","positions; ","")&IF(Deal!$F$43="RED","revenue trend; ","")&IF(Deal!$F$44="RED","neg days/NSFs; ","")&IF(Deal!$F$45="RED","credit; ","")&IF(COUNTIF(Deal!$F$38:$F$45,"KNOCKOUT")>0,"KNOCKOUT present",""))',None,None),
   (33,'Greens at the requested offer','=IF(COUNTIF(Deal!$F$38:$F$45,"GREEN")=0,"none",IF(Deal!$F$38="GREEN","leverage; ","")&IF(Deal!$F$39="GREEN","holdback; ","")&IF(Deal!$F$40="GREEN","daily debt/balance; ","")&IF(Deal!$F$41="GREEN","fund/revenue; ","")&IF(Deal!$F$42="GREEN","positions; ","")&IF(Deal!$F$43="GREEN","revenue trend; ","")&IF(Deal!$F$44="GREEN","clean neg days/NSFs; ","")&IF(Deal!$F$45="GREEN","credit; ",""))',None,None)]
for item in S:
    r,t,v,d,fmt=item
    label(ds,f'B{r}',t,B); put(ds,f'C{r}',v,B if r in (5,6) else N,fmt,align=LEFT if r>=19 else CEN)
    if d is not None: c=ds[f'D{r}']; c.value=d; c.font=NOTE; c.alignment=WRAP; c.border=BOX
    if r>=26: ds.merge_cells(f'C{r}:D{r}'); ds[f'C{r}'].alignment=WRAP
ds['C5'].font=Font(name='Calibri',size=13,bold=True); ds['C6'].font=Font(name='Calibri',size=13,bold=True); ds.row_dimensions[5].height=32; ds.row_dimensions[25].height=40
dec_rules(ds,'C5'); ds.sheet_view.showGridLines=False

# ======================= Comparable Deals =======================
cd=wb.create_sheet('Comparable Deals')
cd['A1']='COMPARABLE DEALS  -  every house-book deal in the primary cohort'; cd['A1'].font=Font(name='Calibri',size=14,bold=True)
cd['A2']=f"=\"Primary cohort: \"&{HS}$C$30&\"   |   \"&{HS}$C$31&\" seasoned / all   |   first 400 listed\""; cd['A2'].font=B
ch=['#','Contract ID','DBA','Status','Funded','Position','Industry','ISO','State','New/renewal','Factor','Advance','Collected','% paid','Net cash','Principal lost','Term (bd)','Age (bd)','Days on book','Seasoned','Default','Open','Similarity tier','Why included']
sub(cd,4,ch); cd.row_dimensions[4].height=30
src={'Contract ID':'A','DBA':'B','Status':'D','Funded':'E','Position':'G','Industry':'Y','ISO':'K','State':'J','New/renewal':'L','Factor':'M','Advance':'N','Collected':'Q','% paid':'S','Net cash':'AJ','Principal lost':'AK','Term (bd)':'AA','Age (bd)':'AB','Days on book':'AN','Seasoned':'AC','Default':'AE','Open':'AH','Similarity tier':'BC'}
for k in range(1,401):
    r=4+k; cd.cell(r,1,k).font=NOTE
    for j,h in enumerate(ch[1:-1]):
        c=cd.cell(r,2+j,f"=IFERROR(INDEX({hr(src[h])},MATCH($A{r},{hr('BE')},0)),\"\")"); c.font=N
        if h in ('Advance','Collected','Net cash','Principal lost'): c.number_format='"$"#,##0'
        if h=='Funded': c.number_format='mm/dd/yyyy'
        if h=='% paid': c.number_format='0%'
        if h=='Factor': c.number_format='0.00'
    c=cd.cell(r,len(ch),f'=IF(B{r}="","",IF(W{r}=1,"NR + industry + position",IF(W{r}=2,"NR + ISO + position",IF(W{r}=3,"NR + industry",IF(W{r}=4,"NR + ISO",IF(W{r}=5,"NR + position",IF(W{r}=6,"NR only","book baseline"))))))&"  ["&IF(T{r}=1,"seasoned","unseasoned")&IF(U{r}=1,", defaulted","")&IF(V{r}=1,", open","")&"]")'); c.font=NOTE
widths(cd,{'A':5,'B':12,'C':30,'D':18,'E':11,'F':8,'G':26,'H':24,'I':6,'J':10,'K':7,'L':11,'M':11,'N':7,'O':11,'P':11,'Q':8,'R':8,'S':8,'T':8,'U':8,'V':7,'W':9,'X':40})
cd.freeze_panes='C5'

# ======================= Checks =======================
ck=wb.create_sheet('Checks'); widths(ck,{'A':3,'B':64,'C':16,'D':70})
ck['B1']='CHECKS / AUDIT'; ck['B1'].font=Font(name='Calibri',size=14,bold=True)
sub(ck,3,['','Check','Result','Detail'],start=1)
checks=[('MISSING INPUTS',None,None),('Bank data entered','=IF(Deal!$C$24=0,"WARN","OK")','="months entered: "&COUNT(Deal!$C$12:$C$23)'),('Average daily balance entered','=IF(Deal!$D$24=0,"WARN","OK")',None),('New deal or renewal entered','=IF(Deal!$K$15="","WARN","OK")','="required for any cohort tighter than the book"'),
 ('ISO entered and found in the house book','=IF(Deal!$K$12="","WARN",IF(COUNTIF('+hr("K")+',Deal!$K$12)=0,"WARN","OK"))',None),('Industry entered and found in the house book','=IF(Deal!$K$13="","WARN",IF(COUNTIF('+hr("Y")+',Deal!$K$13)=0,"WARN","OK"))',None),('Credit score entered','=IF(Deal!$K$16="","WARN","OK")','="blank = not scored"'),('Requested factor is an approved tier','=IF(Deal!$C$6="","OK",IF(ISNUMBER(MATCH(Deal!$C$6,Controls!$B$24:$D$24,0)),"OK","WARN"))',None),
 ('FORMULA ERRORS',None,None),('Deal sheet','=IF(SUMPRODUCT(--ISERROR(Deal!$B$1:$M$92))=0,"OK","FAIL")','=SUMPRODUCT(--ISERROR(Deal!$B$1:$M$92))&" error cells"'),('Historical Data','=IF(SUMPRODUCT(--ISERROR('+hr("Y")+'))+SUMPRODUCT(--ISERROR('+hr("AA")+'))+SUMPRODUCT(--ISERROR('+hr("BC")+'))+SUMPRODUCT(--ISERROR('+hr("BE")+'))=0,"OK","FAIL")',None),('Historical Score','=IF(SUMPRODUCT(--ISERROR(\'Historical Score\'!$A$1:$AA$56))=0,"OK","FAIL")',None),('Offer Engine / Decision Summary','=IF(SUMPRODUCT(--ISERROR(\'Offer Engine\'!$B$1:$G$41))+SUMPRODUCT(--ISERROR(\'Decision Summary\'!$B$1:$D$33))=0,"OK","FAIL")',None),
 ('SAFETY',None,None),('Final fund <= sized max fund',"='Offer Engine'!$E$36",None),('Final daily <= sized max daily',"='Offer Engine'!$E$37",None),('No offer-side metric red/knockout at the final offer',"='Offer Engine'!$E$38",None),('Merchant knockouts untouched',"='Offer Engine'!$E$39",None),('Sizing flex below 100%',"='Offer Engine'!$E$40",None),
 ('COHORT QUALITY',None,None),('Primary cohort met the minimums',"=IF('Historical Score'!$C$29=7,\"WARN\",\"OK\")","=IF('Historical Score'!$C$29=7,\"book baseline - modifier Neutral\",'Historical Score'!$C$30)"),('Seasoned deals behind the primary cohort','=Deal!$F$57','=Deal!$C$57&" seasoned deals"'),('Conflicting ISO vs industry signals',"=IF('Historical Score'!$C$55=\"No\",\"OK\",\"WARN\")","='Historical Score'!$C$55"),('Tighter cohorts skipped for being too small',"=IF('Historical Score'!$C$29=1,\"OK\",\"WARN\")","='Historical Score'!$C$56"),
 ('DATA QUALITY (house book)',None,None),('Deals loaded','="OK"','=COUNTA('+hr("A")+')&" deals"'),('Statuses not in the Controls mapping','=IF(SUMPRODUCT(1-'+hr("AI")+')=0,"OK","FAIL")','=SUMPRODUCT(1-'+hr("AI")+')&" deals"'),('Deals with no revenue figure','=IF(SUMPRODUCT(1-'+hr("AM")+')>0,"WARN","OK")','=SUMPRODUCT(1-'+hr("AM")+')&" deals - excluded from adv/rev averages"'),('Never-paid deals (collected <= 0)','="WARN"','=COUNTIF('+hr("Q")+',"<=0")&" deals"'),('Open status but 99.9%+ paid','=IF(SUMPRODUCT('+hr("AH")+'*('+hr("S")+'>=0.999))=0,"OK","WARN")','=SUMPRODUCT('+hr("AH")+'*('+hr("S")+'>=0.999))&" deals"'),('Open deals older than 1.5x contracted term','="WARN"','=SUMPRODUCT('+hr("AH")+'*('+hr("AB")+'>1.5*'+hr("AA")+'))&" deals"'),('Weekly term above 39 weeks (9-month box)','="WARN"','=SUMPRODUCT(('+hr("T")+'>=12)*('+hr("T")+'>39))&" deals"'),('Payback = advance x factor on every deal','=IF(SUMPRODUCT(--(ABS('+hr("P")+'-'+hr("N")+'*'+hr("M")+')>1))=0,"OK","WARN")','=SUMPRODUCT(--(ABS('+hr("P")+'-'+hr("N")+'*'+hr("M")+')>1))&" mismatches"'),
 ('RECONCILIATION',None,None),('Total funded','="OK"','=TEXT(SUM('+hr("N")+'),"$#,##0")&"  (full book: $34,481,626)"'),('Total collected','="OK"','=TEXT(SUM('+hr("Q")+'),"$#,##0")&"  (full book: $27,438,510)"'),('Principal lost','="OK"','=TEXT(SUMPRODUCT('+hr("AK")+'),"$#,##0")&"  (full book: $5,270,163)"'),('Tier-7 seasoned funded = in-window seasoned funded',"=IF(ABS('Historical Score'!$G$20-SUMPRODUCT("+hr("AP")+"*"+hr("N")+"))<1,\"OK\",\"FAIL\")",None),('Comparable list count = primary cohort matches',"=IF(MIN(400,INDEX('Historical Score'!$D$14:$D$20,'Historical Score'!$C$29))=SUMPRODUCT(--('Comparable Deals'!$B$5:$B$404<>\"\")),\"OK\",\"FAIL\")","=SUMPRODUCT(--('Comparable Deals'!$B$5:$B$404<>\"\"))&\" listed\"")]
r=4
for t,f,d in checks:
    if f is None: hdr(ck,f'B{r}',t,f'B{r}:D{r}')
    else:
        label(ck,f'B{r}',t); put(ck,f'C{r}',f,B)
        if d: c=ck[f'D{r}']; c.value=d; c.font=NOTE; c.alignment=LEFT; c.border=BOX
    r+=1
chk_rules(ck,f'C4:C{r}'); ck['B2']=f'=COUNTIF(C4:C{r},"FAIL")&" FAIL  |  "&COUNTIF(C4:C{r},"WARN")&" WARN  |  "&COUNTIF(C4:C{r},"OK")&" OK"'; ck['B2'].font=B
ck.sheet_view.showGridLines=False

# ======================= README =======================
rd=wb.create_sheet('README',0); widths(rd,{'A':3,'B':125})
txt=["ASPIRE COMBINED UNDERWRITING MODEL  v2"+("  -  DEMO: 8 sample deals, Deal sheet populated" if DEMO else ""),"",
"HOW TO RUN A DEAL","Deal sheet: sections A, B and existing positions as always; DEAL PROFILE (right of the bank data): ISO, industry, new/renewal, credit score. State is optional and only informational. Position comes from C42.",
"Read section C (cash flow), C2 (house-book indicators, same colour code), D (decision on the requested offer), E (max offer at the sizing line), F (final offer). Decision Summary has the full read-out; Comparable Deals lists the deals behind it.",
"","HOW THE HOUSE BOOK CHANGES THE OFFER","1. Historical Score matches the deal to seven cohorts (new/renewal + industry + position ... down to the whole seasoned book) and picks the first one with enough seasoned deals and dollars.",
"2. Its principal lost rate is blended toward the book by credibility Z = n/(n+K) and divided by the book rate -> relative index (1.00 = book).",
"3. The index lands on a row of the modifier table (Controls section 5). Supportive book -> SIZING FLEX moves the sizing lines toward the red lines (recommended 90% for Strong positive, 45% for Mild positive), so the max offer is larger and metrics may show YELLOW. Neutral -> flex 0 = the tool exactly as built. Negative -> green lines plus a fund multiplier, shorter band, higher factor tier, weekly.",
"4. Section E sizes the max offer at the sizing lines (column H of section C). Offer Engine shows green-line base vs sized vs final side by side and re-checks every metric at the final offer: yellow is allowed, red never.",
"","HARD RULES","History never removes a knockout, never turns a metric red (flex is capped below 100%), never exceeds the sized max, and is Neutral whenever no cohort meets the minimums.",
"","COLOUR KEY","Blue on yellow = input. Green font = link to Controls. Purple bold (column H, section C) = sizing line in use. Same GREEN / YELLOW / RED cells in C2 as in C."]
for i,t in enumerate(txt):
    c=rd.cell(2+i,2,t); c.font=B if (t.isupper() or i==0) else N; c.alignment=Alignment(wrap_text=True,vertical='top')
rd.sheet_view.showGridLines=False

# ---------- demo population ----------
if DEMO:
    d=ws
    d['B6']=40000; d['C6']=1.45; d['D6']=24; d['E6']='Weekly'
    for i,(m,rev,bal,neg,nsf) in enumerate([('Aug',88000,7400,0,0),('Jul',84000,6900,0,1),('Jun',91000,8100,0,0),('May',82000,6500,1,0)]):
        r=12+i; d[f'B{r}']=m; d[f'C{r}']=rev; d[f'D{r}']=bal; d[f'E{r}']=neg; d[f'F{r}']=nsf
    d['K12']='Fundwell'; d['K13']='Eating & Drinking Places'; d['K14']='NJ'; d['K15']='New Deal'; d['K16']=672

for shn in wb.sheetnames:
    sh=wb[shn]; sh.sheet_properties.pageSetUpPr.fitToPage=True; sh.page_setup.fitToWidth=1; sh.page_setup.fitToHeight=0
wb._sheets=[wb[n] for n in ['README','Deal','Decision Summary','Offer Engine','Historical Score','Comparable Deals','Historical Data','Controls','Checks','Ref']]
wb.active=1; wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print('saved',OUT,'rows',NR-1)
