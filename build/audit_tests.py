import openpyxl, os, shutil, subprocess, math
from openpyxl.utils import get_column_letter
SRC="/root/.claude/uploads/15c93af4-845a-5d84-bd2c-2f824a3823d2/d23c6065-Underwriting_Tool_Combined_with_House_book_analytics_1.xlsx"
W="/tmp/claude-0/-home-user-Underwriting-Tool-House-Book/15c93af4-845a-5d84-bd2c-2f824a3823d2/scratchpad/audit"
shutil.rmtree(W,ignore_errors=True); os.makedirs(W,exist_ok=True)

def recalc(path,tag):
    od=os.path.join(W,tag); os.makedirs(od,exist_ok=True)
    subprocess.run(["libreoffice","--headless","--convert-to","xlsx","--outdir",od,path],
                   check=True,capture_output=True,timeout=1800)
    return openpyxl.load_workbook(os.path.join(od,os.path.basename(path)),data_only=True)

def run(tag, edits):
    wb=openpyxl.load_workbook(SRC)
    # LibreOffice can't do XLOOKUP - swap for the test only
    wb["Primary Combination Calc"]["C36"]='=IF(MAX($I$23:$I$26)=0,"none rated",INDEX($B$23:$B$26,MATCH(MAX($I$23:$I$26),$I$23:$I$26,0)))'
    for sheet,cell,val in edits: wb[sheet][cell]=val
    f=os.path.join(W,f"{tag}.xlsx"); wb.save(f)
    v=recalc(f,tag+"o"); d=v["Deal"]; pc=v["Primary Combination Calc"]
    return dict(dec=d["C75"].value, rec=d["C77"].value, cf=d["C76"].value,
                term=d["C71"].value, fac=d["D71"].value, daily=d["C69"].value,
                score=d["C65"].value, reds=d["C60"].value, ko=d["E60"].value,
                yel=sum(1 for c in ("C86","C87","C88","C89") if d[c].value=="YELLOW")
                    +sum(1 for r in range(38,46) if d[f"F{r}"].value=="YELLOW"),
                cohort=pc["C29"].value, tier=pc["D29"].value, idx=pc["C33"].value,
                out=pc["C34"].value, flex=pc["C35"].value, ev=pc["C37"].value)

TESTS=[
 ("T1  baseline (as shipped)", []),
 ("T2  blank deal profile",    [("Deal","K12",None),("Deal","K13",None),("Deal","K14",None),("Deal","K15",None)]),
 ("T3  blank bank data",       [("Deal",f"{c}{r}",None) for r in range(12,24) for c in "CDEF"]),
 ("T4  neutral book (flex 0)", [("Controls","B27",-1),("Controls","B28",99)]),
 ("T5  negative book",         [("Controls","B27",-1),("Controls","B28",0.05)]),
 ("T6  KNOCKOUT on row 42 pos",[("Deal","C28",100000),("Deal","D28",5000),
                                ("Deal","C29",1),("Deal","D29",1),("Deal","C30",1),("Deal","D30",1),
                                ("Deal","C31",1),("Deal","D31",1),("Deal","C32",1),("Deal","D32",1),
                                ("Deal","C33",1),("Deal","D33",1)]),
 ("T7  KNOCKOUT on row 38 LVG",[("Deal","D28",30000)]),
 ("T8  one RED (credit 640)",  [("Deal","K15",640)]),
 ("T9  two REDs (cr+trend)",   [("Deal","K15",620),("Deal","C13",70000),("Deal","C14",100000)]),
 ("T10 tiny revenue",          [("Deal",f"C{r}",9000) for r in range(12,16)]),
]
print(f"{'test':<30}{'DECISION':<13}{'rec fund':>10}{'cashflow':>10}{'term':>6}{'fac':>6}{'score':>6}{'R':>3}{'KO':>3}{'Y':>3}  {'tier':<5}{'outcome':<12}{'idx':>7}")
print("-"*140)
res={}
for tag,ed in TESTS:
    try:
        r=run(tag.split()[0],ed); res[tag]=r
        fmt=lambda x,f="{:,.0f}": (f.format(x) if isinstance(x,(int,float)) else str(x)[:9])
        print(f"{tag:<30}{str(r['dec']):<13}{fmt(r['rec']):>10}{fmt(r['cf']):>10}{fmt(r['term']):>6}"
              f"{fmt(r['fac'],'{:.2f}'):>6}{fmt(r['score'],'{:.1f}'):>6}{fmt(r['reds']):>3}{fmt(r['ko']):>3}{r['yel']:>3}  "
              f"{str(r['tier']):<5}{str(r['out'])[:11]:<12}{fmt(r['idx'],'{:.3f}'):>7}")
    except Exception as e:
        print(f"{tag:<30}ERROR {e}")
