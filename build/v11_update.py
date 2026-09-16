import openpyxl, datetime as dt
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter as L

SRC='/root/.claude/uploads/15c93af4-845a-5d84-bd2c-2f824a3823d2/c314f08b-Underwriting_Tool_Combined_with_House_book_analytics_4.xlsx'
OUT='/home/user/Underwriting-Tool-House-Book/Aspire_v11_RENEWAL_FIX_highlighted.xlsx'
BLUE=PatternFill('solid', fgColor='FFCCE5FF')
LAST=4603
wb=openpyxl.load_workbook(SRC)
h=wb['Historical Data']; c=wb['Controls']; cd=wb['Comparable Deals']
touched=0
def mark(ws,cell):
    global touched
    ws[cell].fill=BLUE; touched+=1

# ---- 1. TRIM legal name (col C) so renewal matching is reliable
for r in range(2,LAST+1):
    v=h.cell(r,3).value
    if isinstance(v,str):
        s=v.strip()
        if s!=v:
            h.cell(r,3).value=s
            mark(h,f'C{r}')

# ---- 2. Renewal flag (col L) computed as values
names={}; dates={}
for r in range(2,LAST+1):
    n=h.cell(r,3).value; e=h.cell(r,5).value
    n=n.strip().upper() if isinstance(n,str) else None
    names[r]=n
    dates[r]=e.date() if isinstance(e,dt.datetime) else e
by={}
for r,n in names.items():
    if n: by.setdefault(n,[]).append(r)
ren=0
for r in range(2,LAST+1):
    n=names[r]; e=dates[r]
    if not n:
        val=''
    else:
        val='New Deal'
        if e is not None:
            for o in by[n]:
                if o!=r and dates[o] is not None and dates[o]<e:
                    val='Renewal'; break
    if val=='Renewal': ren+=1
    h.cell(r,12).value=val
    mark(h,f'L{r}')
print('renewals written:',ren)

# ---- 3. Controls: add Cancelled to the status table
c['A45']='Cancelled'; c['B45']='No'; c['C45']='Yes'; c['D45']='Yes'; c['E45']='No'
c['F45']=f"=COUNTIF('Historical Data'!$D$2:$D${LAST},A45)"
for col in 'ABCDEF': mark(c,f'{col}45')

# ---- 4. Widen the status lookups from row 44 -> 45  (AE AF AG AH AI)
for r in range(2,LAST+1):
    for col,idx in ((31,2),(32,3),(33,4),(34,5)):
        h.cell(r,col).value=f'=IF(IFERROR(VLOOKUP(D{r},Controls!$A$34:$E$45,{idx},FALSE()),"No")="Yes",1,0)'
        mark(h,f'{L(col)}{r}')
    h.cell(r,35).value=f'=IF(ISNUMBER(MATCH(D{r},Controls!$A$34:$A$45,0)),1,0)'
    mark(h,f'AI{r}')

# ---- 5. Comparable Deals: extend the mirror to the full book
SRCCOL=['A','B','D','E','G','Y','K','L','M','N','Q','S','AJ','AK','AA','AB','AN','AC','AE','AH','AU']
for i,sc in enumerate(SRCCOL,start=1):
    for r in range(5,5+(LAST-1)):
        hd=r-3
        cd.cell(r,i).value=f"='Historical Data'!{sc}{hd}"
        if r>1071: mark(cd,f'{L(i)}{r}')

wb.save(OUT)
print('cells highlighted:',touched)
print('saved',OUT)
