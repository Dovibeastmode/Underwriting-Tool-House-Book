import openpyxl, re, datetime
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.formula import ArrayFormula

SRC="/root/.claude/uploads/15c93af4-845a-5d84-bd2c-2f824a3823d2/f47874f8-House_Book_Analytics_3.xlsx"
OUT="House_Book_Analytics_v4_SEASONING.xlsx"
LAST_ROW=1068          # real last DEAL row (1069+ is the column-documentation table)
OLD_LAST=1068
BLUE=PatternFill("solid",fgColor="FFCCE5FF")
def t(v): return v.text if isinstance(v,ArrayFormula) else v

wb=openpyxl.load_workbook(SRC)
ct=wb["Controls"]; d=wb["Data"]
touched=[]

# ---------- 1. Controls: the seasoning block ----------
def put(ws,coord,val,bold=False,note=False):
    ws[coord]=val
    ws[coord].fill=BLUE
    if bold: ws[coord].font=Font(bold=True)
    if note: ws[coord].font=Font(italic=True,size=9,color="FF555555")
    touched.append(f"{ws.title}!{coord}")

put(ct,"A8","SEASONING   -  a deal only counts as evidence once it has had time to fail",bold=True)
put(ct,"A9","As-of date (when this data was pulled)")
put(ct,"B9",datetime.datetime(2026,8,19))
ct["B9"].number_format="yyyy-mm-dd"
put(ct,"A10","Seasoning multiple  (age >= multiple x expected term)")
put(ct,"B10",1.25)
put(ct,"D9","Require seasoning?  Y / N")
put(ct,"E9","Y")
put(ct,"D10","Seasoned deals in window")
put(ct,f"E10",f"=SUMPRODUCT(Data!$BE$2:$BE${LAST_ROW}*Data!$AK$2:$AK${LAST_ROW})")
put(ct,"F10","of "+str(LAST_ROW-1)+" deals",note=True)

# ---------- 2. Data: three new columns ----------
d["BC1"]="Age at as-of (business days)"; d["BC1"].fill=BLUE
d["BD1"]="Terminal outcome?";            d["BD1"].fill=BLUE
d["BE1"]="Seasoned";                     d["BE1"].fill=BLUE
for r in range(2,LAST_ROW+1):
    d[f"BC{r}"]=f'=IF($E{r}="","",NETWORKDAYS($E{r},Controls!$B$9,Ref!$F$4:$F$47))'
    d[f"BD{r}"]=f'=IF(OR($AH{r}=1,$AJ{r}=1),1,0)'
    d[f"BE{r}"]=(f'=IF(OR($BD{r}=1,AND($AE{r}<>"",$BC{r}>=Controls!$B$10*$AE{r})),1,0)')
    for c in ("BC","BD","BE"): d[f"{c}{r}"].fill=BLUE
touched.append(f"Data!BC2:BE{LAST_ROW}")

# ---------- 3. In scope now carries seasoning ----------
for r in range(2,LAST_ROW+1):
    d[f"AK{r}"]=(f'=IF(AND($E{r}>=Controls!$B$5,$E{r}<=Controls!$B$6,'
                 f'OR(Controls!$E$9<>"Y",$BE{r}=1)),1,0)')
    d[f"AK{r}"].fill=BLUE
touched.append(f"Data!AK2:AK{LAST_ROW}")

# (no range change - $2:$1068 was already correct; rows 1069+ are documentation)

wb.save(OUT)
print("saved",OUT)
print("controls cells written:", [x for x in touched if x.startswith("Controls")])
