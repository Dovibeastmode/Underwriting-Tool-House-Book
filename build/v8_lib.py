import openpyxl, subprocess, os, shutil, glob
SRC="/root/.claude/uploads/15c93af4-845a-5d84-bd2c-2f824a3823d2/ae963fcd-Aspire_Combined_Underwriting_Model_v7_FULL_1.xlsx"
LAST=1068
HD="'Historical Data'!"

def unmerge(ws, coord):
    for rng in list(ws.merged_cells.ranges):
        if coord in rng: ws.unmerge_cells(str(rng))

def restore_baseline(ws_hs):
    """v7 fixes their file is missing, so old and new start from the same place."""
    ws_hs["D29"]='=IF($C$5="","",IFERROR(MATCH("Yes",$E$14:$E$21,0),8))'
    ws_hs["E29"]='row # (read by Historical Data)'
    for c,f in [("C29",'=IF($D$29="","- no deal entered -",INDEX($B$14:$B$21,$D$29))'),
                ("C30",'=IF(OR($D$29="",$D$29=8),"",INDEX($C$14:$C$21,$D$29))'),
                ("C31",'=IF(OR($D$29="",$D$29=8),"",INDEX($F$14:$F$21,$D$29))'),
                ("C32",'=IF(OR($D$29="",$D$29=8),"",INDEX($G$14:$G$21,$D$29))'),
                ("C33",'=IF(OR($D$29="",$D$29=8),"",INDEX($H$14:$H$21,$D$29))')]:
        ws_hs[c]=f
    # LibreOffice cannot evaluate XLOOKUP; swap to INDEX/MATCH for verification only
    ws_hs["C36"]='=IF(MAX($I$23:$I$26)=0,"none rated",INDEX($B$23:$B$26,MATCH(MAX($I$23:$I$26),$I$23:$I$26,0)))'

# ---------- the v8 cohort engine ----------
# flag columns on Historical Score:  K=req NR  L=req ISO  M=req Ind  N=req Pos
DIMS=[("K","AR"),("L","AT"),("M","AS"),("N","AQ")]   # (flag col, Historical Data match col)
TIER_FLAGS={14:(1,1,1,0), 15:(1,0,1,1), 16:(1,1,0,1), 17:(1,0,1,0),
            18:(1,1,0,0), 19:(1,0,0,1), 20:(1,0,0,0), 21:(0,0,0,0)}

def mask(r):
    """m*req + (1-req):  req=1 -> must match, req=0 -> ignore."""
    return "*".join(f"({HD}${m}$2:${m}${LAST}*${f}{r}+(1-${f}{r}))" for f,m in DIMS)

def apply_v8(wb):
    hs=wb["Historical Score"]; hd=wb["Historical Data"]; dl=wb["Deal"]

    # 1. tier flag table
    for col,hdr in zip("KLMN",["req NR","req ISO","req Ind","req Pos"]):
        hs[f"{col}13"]=hdr
    hs["O13"]="<- tier definition: 1 = dimension must match, 0 = ignore it"
    for r,flags in TIER_FLAGS.items():
        for col,v in zip("KLMN",flags): hs[f"{col}{r}"]=v

    # 2. one generic formula per stat, identical in every tier row
    for r in TIER_FLAGS:
        M=mask(r)
        hs[f"C{r}"]=f"=SUMPRODUCT({HD}$AP$2:$AP${LAST}*{M})"
        hs[f"D{r}"]=f"=SUMPRODUCT({HD}$AP$2:$AP${LAST}*{M}*{HD}$N$2:$N${LAST})"
        hs[f"F{r}"]=f'=IF(D{r}=0,"",SUMPRODUCT({HD}$AP$2:$AP${LAST}*{M}*{HD}$AK$2:$AK${LAST})/D{r})'
        hs[f"G{r}"]=f'=IF(D{r}=0,"",SUMPRODUCT({HD}$AP$2:$AP${LAST}*{M}*{HD}$AJ$2:$AJ${LAST})/D{r})'

    # 3. selection - whole-book row always qualifies, so MATCH can never fail
    hs["D29"]='=IF($C$5="","",MATCH("Yes",$E$14:$E$21,0))'
    hs["E29"]='row # (read by Historical Data)'
    for col,k in zip("KLMN",range(1,5)):
        hs[f"{col}29"]=f'=IF($D$29="","",INDEX({col}$14:{col}$21,$D$29))'
    hs["O29"]="<- flags of the selected tier (read by Historical Data BA)"

    # 4. one cell owns the whole no-evidence question
    hs["D30"]='=COUNTA($B$14:$B$21)'; hs["E30"]='fallback row # (whole seasoned book)'
    hs["B37"]='HAS EVIDENCE?'
    hs["C37"]='=IF(OR($D$29="",$D$29=$D$30),"NO","YES")'
    hs["D37"]='<- every guard downstream reads this, not a hard-coded row number'

    # 5. guards now read C37
    unmerge(hs,"C29")
    hs["C29"]='=IF($D$29="","- no deal entered -",INDEX($B$14:$B$21,$D$29))'
    for c,src in [("C30","$C$14:$C$21"),("C31","$F$14:$F$21"),
                  ("C32","$G$14:$G$21"),("C33","$H$14:$H$21")]:
        unmerge(hs,c)
        hs[c]=f'=IF($C$37="NO","",INDEX({src},$D$29))'
    hs["C34"]=('=IF(OR($C$37="NO",$C$33=""),"Neutral",'
               'IF($C$33<=Controls!$B$29,"Supportive",'
               'IF($C$33>Controls!$B$30,"Negative","Neutral")))')

    # 6. retire the hardcoded combo columns
    for col in ["AU","AV","AW","AX","AY","AZ"]:
        hd[f"{col}1"]=None
        for r in range(2,LAST+1): hd[f"{col}{r}"]=None
    hd["AU1"]="(retired - combos are now defined by the flag table on Historical Score)"

    # 7. BA reads the selected tier's flags - same mask, no CHOOSE
    for r in range(2,LAST+1):
        prod="*".join(f"({m}{r}*'Historical Score'!${f}$29+1-'Historical Score'!${f}$29)"
                      for f,m in DIMS)
        hd[f"BA{r}"]=f'=IF(OR(AD{r}=0,\'Historical Score\'!$D$29=""),0,{prod})'

    # 8. Deal guards
    dl["F51"]=('=IF(\'Historical Score\'!$C$37="NO","NO DEALS",'
               'IF(\'Historical Score\'!$C$33<=Controls!$B$29,"GREEN",'
               'IF(\'Historical Score\'!$C$33>Controls!$B$30,"RED","YELLOW")))')
    dl["C56"]='=IF(\'Historical Score\'!$C$37="NO","",\'Historical Score\'!$C$30)'

# ---------- harness ----------
PROFILES=[
 ("their example",      "New Deal", 1, "Agricultural Services",     "United Secured Capital 8"),
 ("common new deal",    "New Deal", 3, "Eating & Drinking Places",  "GREYSTONE"),
 ("renewal",            "Renewal",  2, "Trucking & Warehousing",    "BFA 3"),
 ("thin ISO+industry",  "New Deal", 1, "Nonmetallic Minerals Mining","Fidelity Funding"),
 ("no ISO/industry",    "New Deal", 4, "",                          ""),
 ("unmatchable",        "New Deal", 99,"ZZZ Nonexistent",           "ZZZ Nonexistent"),
 ("blank deal",         "",         "", "",                         ""),
]

def recalc(path, outdir):
    os.makedirs(outdir,exist_ok=True)
    subprocess.run(["libreoffice","--headless","--convert-to","xlsx","--outdir",outdir,path],
                   check=True, capture_output=True, timeout=600)
    return os.path.join(outdir, os.path.basename(path))

def probe(base_path, tag, mutate, workdir):
    """Write one file per profile with C5:C8 pinned, recalc, read the engine's output."""
    out={}
    for i,(name,nr,pos,ind,iso) in enumerate(PROFILES):
        wb=openpyxl.load_workbook(base_path)
        mutate(wb)
        hs=wb["Historical Score"]
        for c,v in [("C5",nr),("C6",pos),("C7",ind),("C8",iso)]:
            hs[c]= v if v!="" else None
        f=os.path.join(workdir,f"{tag}_{i}.xlsx")
        wb.save(f)
        r=recalc(f, os.path.join(workdir,f"{tag}_{i}_out"))
        v=openpyxl.load_workbook(r,data_only=True)
        h=v["Historical Score"]
        out[name]=dict(
            D29=h["D29"].value, cohort=h["C29"].value, n=h["C30"].value,
            plr=h["C31"].value, rof=h["C32"].value, idx=h["C33"].value,
            outcome=h["C34"].value, flex=h["C35"].value,
            tiers=[(h[f"C{r}"].value, h[f"D{r}"].value, h[f"E{r}"].value,
                    h[f"F{r}"].value, h[f"H{r}"].value) for r in range(14,22)],
            BAsum=sum((v["Historical Data"][f"BA{rr}"].value or 0) for rr in range(2,LAST+1)),
        )
    return out
