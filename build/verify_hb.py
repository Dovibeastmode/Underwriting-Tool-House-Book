import openpyxl, os, shutil, subprocess
W="/tmp/claude-0/-home-user-Underwriting-Tool-House-Book/15c93af4-845a-5d84-bd2c-2f824a3823d2/scratchpad/hb"
shutil.rmtree(W,ignore_errors=True); os.makedirs(W,exist_ok=True)
SRC="House_Book_Analytics_v4_SEASONING.xlsx"
ORIG="/root/.claude/uploads/15c93af4-845a-5d84-bd2c-2f824a3823d2/f47874f8-House_Book_Analytics_3.xlsx"

def recalc(path,tag):
    od=os.path.join(W,tag); os.makedirs(od,exist_ok=True)
    subprocess.run(["libreoffice","--headless","--convert-to","xlsx","--outdir",od,path],
                   check=True,capture_output=True,timeout=1800)
    return openpyxl.load_workbook(os.path.join(od,os.path.basename(path)),data_only=True)

def variant(tag, toggle):
    wb=openpyxl.load_workbook(SRC)
    if toggle is not None: wb["Controls"]["E9"]=toggle
    f=os.path.join(W,f"{tag}.xlsx"); wb.save(f)
    return recalc(f,tag+"o")

def stats(v,label):
    d=v["Data"]
    rows=range(2,1106)
    inscope=sum(1 for r in rows if d.cell(r,37).value==1)          # AK col 37
    seas=sum(1 for r in rows if d.cell(r,57).value==1)             # BE col 57
    fu=sum(d.cell(r,16).value or 0 for r in rows if d.cell(r,37).value==1)   # P advance
    co=sum(d.cell(r,19).value or 0 for r in rows if d.cell(r,37).value==1)   # S collected
    lo=sum(d.cell(r,40).value or 0 for r in rows if d.cell(r,37).value==1)   # AN principal lost
    df=sum(d.cell(r,34).value or 0 for r in rows if d.cell(r,37).value==1)   # AH default
    print(f"  {label:<26} in-scope {inscope:<5} seasoned {seas:<5} funded ${fu:>12,.0f} "
          f"PLR {lo/fu:>6.1%}  ROF {(co-fu)/fu:>+6.1%}  dflt {df/inscope:>5.1%}")

print("recalculating (this takes a few minutes)...")
o=recalc(ORIG,"orig")
do=o["Data"]; rows=range(2,1106)
inscope=sum(1 for r in rows if do.cell(r,37).value==1)
fu=sum(do.cell(r,16).value or 0 for r in rows if do.cell(r,37).value==1)
co=sum(do.cell(r,19).value or 0 for r in rows if do.cell(r,37).value==1)
lo=sum(do.cell(r,40).value or 0 for r in rows if do.cell(r,37).value==1)
df=sum(do.cell(r,34).value or 0 for r in rows if do.cell(r,37).value==1)
print(f"\n  {'ORIGINAL (no seasoning)':<26} in-scope {inscope:<5} seasoned {'-':<5} funded ${fu:>12,.0f} "
      f"PLR {lo/fu:>6.1%}  ROF {(co-fu)/fu:>+6.1%}  dflt {df/inscope:>5.1%}")
stats(variant("off","N"), "v4 toggle = N")
stats(variant("on","Y"),  "v4 toggle = Y")
